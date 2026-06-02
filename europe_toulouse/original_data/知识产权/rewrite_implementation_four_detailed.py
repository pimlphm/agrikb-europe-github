from __future__ import annotations

import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement


IP = Path(__file__).resolve().parent
SRC = IP / "发明专利_重构申请文档_知产权综合润色版_附图硬件增强版.docx"
OUT = IP / "发明专利_重构申请文档_知产权综合润色版_附图硬件增强版_第四部分详化版.docx"
REPORT = IP / "专利文档_附图硬件增强_修改报告.md"


def find_para(doc: Document, prefix: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i
    raise RuntimeError(f"marker not found: {prefix}")


def remove_para(p):
    p._element.getparent().remove(p._element)


def add_after(anchor, text: str, bold: bool = False):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    p = anchor._parent.add_paragraph()
    p._p = new_p
    r = p.add_run(text)
    r.font.name = "宋体"
    r.font.size = Pt(10.5)
    r.bold = bold
    return p


IMPLEMENTATION = [
    ("导言", [
        "以下结合附图和实施例，对本发明作进一步详细说明。以下四个实施例按照工程落地顺序组织，分别覆盖系统部署与数据对象、资料解析入图与校验、航空维修PDF端到端应用、国产计算卡高吞吐推理部署。每个实施例均以可实施为原则，说明输入、输出、对象字段、处理步骤、判断条件、异常回退和审计记录。所属领域技术人员可依据下列步骤实现等同系统，而不受具体软件库、模型名称、硬件厂商或界面形态限制。",
        "在以下实施例中，“正式对象”是指通过格式、本体、标准引用、证据跨度、冲突和置信度校验后写入正式知识图谱或正式经验库的对象；“候选对象”是指存在来源证据但尚未满足全部校验条件，需要人工复核或补充证据的对象；“证据包”是指围绕一次抽取、检索或问答任务生成的EvidenceBundle，其至少包含证据编号、来源文档、页码、段落、摘要、关联实体、关联标准、置信度和审计编号。",
    ]),
    ("实施例1：系统部署、基础数据对象和硬件适配初始化", [
        "[0030] 本实施例说明系统从空白部署环境到可接收航空维修资料、可选择推理后端并可写入审计日志的初始化过程。系统至少包括资料接入模块、文档解析模块、语义分块模块、本体知识图谱模块、标准条款检索模块、多角色模型推理模块、证据锚定生成模块、经验库模块、审计日志模块、图形复核模块和硬件适配模块。上述模块可部署在同一工作站、内网服务器或私有化服务集群中，也可按检索服务、图谱服务和生成服务拆分部署。",
        "[0031] 部署前准备包括：在目标设备上创建应用根目录，配置docs、data/raw、data/processed、knowledge/chunks、ontology、standards、experience、output/reports、output/graphs和audit目录；为每个目录设置读写权限；准备config.yaml或等效配置对象；准备标准库文件、实体类型表、关系类型表、提示词模板、权限策略和默认阈值表。若系统部署在涉密或内网环境中，上述目录应位于内网存储或本地磁盘，不依赖外部云端路径。",
        "[0032] 配置文件至少包括以下字段：backend表示模型推理后端；model_name表示默认生成模型；embedding_backend表示向量化后端；max_context_tokens表示最大上下文长度；chunk_size和chunk_overlap表示分块长度和重叠长度；retrieval.top_k表示召回数量；retrieval.sparse_weight和retrieval.dense_weight表示稀疏与稠密检索权重；rerank_top_n表示重排候选数量；audit.enabled表示是否写入审计；hardware.priority表示CPU、GPU、NPU、DCU、MLU等设备优先级；fallback.max_concurrency表示回退模式最大并发数。",
        "[0033] 初始化步骤包括：S101，读取配置文件，若缺失必要字段则使用默认值并写入配置缺省日志；S102，创建raw、processed、chunk、graph、standard、experience、audit和export目录；S103，加载标准条款库、实体类型表、关系类型表、关联强度权重表、提示词模板和权限策略；S104，执行硬件探测，形成DeviceProfile对象；S105，启动检索索引、图谱存储和模型服务健康检查；S106，向审计日志写入system_boot事件。",
        "[0034] DeviceProfile对象至少包括device_type、vendor、device_id、memory_total、memory_free、driver_version、runtime_version、backend_name、supported_precision、supported_attention_backend、tensor_parallel_capability、streaming_supported、health_status和fallback_priority字段。若未检测到可用加速卡，device_type被设为CPU，系统自动启用低并发、短上下文和保守生成策略；若检测到多个设备，系统按设备健康状态、剩余内存和配置优先级选择生成后端、向量化后端和重排序后端。",
        "[0035] 硬件探测的具体判定规则为：若设备查询命令或驱动接口返回可用设备且memory_free大于预设阈值，则health_status设为available；若设备存在但驱动版本不匹配、算子库不可用或健康检查失败，则health_status设为degraded；若设备不可见或访问权限不足，则health_status设为unavailable。对于degraded设备，系统只允许执行向量化、重排序或低风险生成任务；对于unavailable设备，系统不分配任务，并将原因写入AuditRecord。",
        "[0036] 系统为每个输入文件创建DocumentRecord对象，字段包括document_id、file_name、file_type、source_path、file_hash、version_tag、security_level、import_time、parser_status、page_count、language、domain_tags、owner和retention_policy。document_id可由文件指纹、导入时间和版本号组合生成。若同一文件重复导入且file_hash一致，系统不重复解析，而是复用已存在的DocumentRecord；若文件名相同但file_hash不同，则创建新版本并保留旧版本引用。",
        "[0037] 系统的核心中间对象包括ChunkRecord、EntityRecord、RelationRecord、StandardClause、EvidenceRecord、ExperienceEntry、QueryRecord、DeviceProfile和AuditRecord。ChunkRecord保存chunk_id、document_id、page_no、section_title、chunk_type、text、table_meta、source_span和embedding_status；EntityRecord保存entity_id、name、normalized_name、entity_type、ontology_layer、properties、source_span、standard_refs和confidence；RelationRecord保存relation_id、head_id、tail_id、relation_type、strength、color_level、evidence_ids、standard_refs、direction和audit_status。",
        "[0038] EvidenceRecord保存evidence_id、document_id、page_no、quote_summary、source_span、standard_refs、related_entities、reliability和created_by；ExperienceEntry保存experience_id、aircraft_type、task_type、applicable_conditions、content、evidence_ids、standard_refs、review_status和usage_count；QueryRecord保存query_id、query_text、query_type、route、evidence_bundle_id、answer_id和status；AuditRecord保存operation_id、operator、time、model_version、hardware_backend、parameters、input_summary、output_summary、latency、token_stat和fallback_event。上述对象统一通过document_id、source_span、evidence_id和audit_id互相引用，不以本地文件路径作为业务主键。",
        "[0039] 初始化完成的判定条件为：标准库可查询，资料目录可写入，图谱存储可创建节点和边，检索索引可接受新增chunk，模型后端返回健康状态或启用离线回退模式，审计日志可追加写入。若任一条件失败，系统不进入正式知识抽取流程，而是输出初始化异常报告，列出失败模块、失败原因、回退策略和人工处理建议。初始化异常报告至少包括error_code、failed_module、detected_state、fallback_action和manual_action字段。",
        "[0040] 本实施例的输出包括系统运行目录、配置快照、标准库加载结果、硬件探测结果、空图谱存储、空检索索引、审计启动记录和初始化报告。该输出可作为后续资料解析、图谱构建、问答生成和硬件推理调度的共同基础。"
    ]),
    ("实施例2：资料解析入图、标准注入、多角色协同和六层校验", [
        "[0041] 本实施例说明如何把PDF、Word、Markdown、文本和表格资料转换为可审计知识图谱。资料接入模块接收输入文件后，根据文件类型选择解析器；对于PDF，提取页码、正文、表格、标题层级、页眉页脚和图片说明；对于Word或Markdown，提取段落、标题、表格和列表；对于扫描质量较差的页面，标记为需要OCR或人工复核。解析结果统一写入PageRecord和BlockRecord。",
        "[0042] PDF解析步骤包括：S201，读取文档元数据和页数；S202，逐页提取文本行、表格区域和图片说明；S203，删除页眉页脚、页码和重复水印；S204，保留章节标题、表格列名、项目符号和编号层级；S205，对每一页计算text_quality，若字符过少、乱码比例过高或表格结构破碎，则将该页标记为low_quality_page；S206，输出PageRecord。PageRecord至少包括document_id、page_no、raw_text、clean_text、tables、figures、quality_score和warnings。",
        "[0043] 语义分块步骤包括：S211，按章节标题、页码、表格边界和段落长度生成初始块；S212，识别块类型，至少包括故障现象块、维修动作块、标准条款块、培训要求块、安全评估块、追溯证据块和一般说明块；S213，对超过长度阈值的块按句子边界继续拆分，对过短且同属一页一节的块进行合并；S214，为每个块生成source_span，记录document_id、page_no、section_title、paragraph_index、char_start和char_end；S215，写入ChunkRecord并进入候选抽取队列。",
        "[0044] chunk_type的判定规则为：含“故障、告警、失效、异常、灯亮、压力低”等词的块优先判为故障现象块；含“检查、更换、复位、记录、签署、放行、复查”等词的块优先判为维修动作块；含“CCAR、FAR、CS、MEL、AMM、标准、规范、条款”等词的块优先判为标准条款块；含“培训、实作、考核、训练、注意事项”等词的块优先判为培训要求块；若同一块满足多种条件，则保留多个domain_tags，并在后续抽取中分别处理。",
        "[0045] 本体约束抽取按照三层本体执行。顶层类型包括实体、过程、功能、质量、角色和倾向；核心层类型包括需求、设计、验证、接口、约束、规范、组件、系统、测试、风险、标准、文档、计划、报告、决策、指标、工具和人员；航空维修领域层类型至少包括故障现象、失效状态、机载系统、部件、维修动作、MEL条目、放行限制、培训科目、人为因素、维修记录、安全目标、DAL分配和证据片段。抽取时，模型或规则引擎只能从允许类型中选择实体类型，不能自由生成未登记类型。",
        "[0046] 实体抽取流程为：S221，从ChunkRecord中识别候选术语；S222，使用术语词典、标准编号规则、英文缩写规则和上下文触发词判断候选类型；S223，将候选实体写成EntityRecord；S224，若同一document_id和相近source_span内出现同名实体，则合并为同一局部实体；S225，若跨文档出现相似实体，则暂存为candidate_same_as关系，等待跨文档共指消解确认。实体名称归一化时，系统统一大小写、去除无意义空格、扩展常见英文缩写，并保留原始名称作为alias。",
        "[0047] 关系抽取流程为：S231，根据关系触发词、句法方向、表格列名和章节语义生成候选关系；S232，根据关系类型允许的头尾实体类型过滤不合法关系；S233，计算关联强度；S234，依据强度区间映射为七级显示等级；S235，将通过阈值的关系写入RelationRecord，将未通过阈值但证据存在的关系写入候选池。关联强度S由关系类型权重、证据距离权重、标准条款权重、图结构权重和模型置信度加权得到，任一分量缺失时使用0值并写入strength_warning字段。",
        "[0048] 标准条款检索采用稀疏检索、稠密检索和类型匹配三路融合。系统从当前ChunkRecord和EntityRecord中抽取故障码、系统名、部件名、MEL编号、标准编号和动作词；BM25路径检索标准编号、中文标题、英文标题和关键词；向量路径检索语义相近条款；类型匹配路径判断条款适用实体类型是否覆盖当前实体类型；最后按R(q,c)=λ1·BM25+λ2·CosSim+λ3·TypeMatch+λ4·ProjectContext计算得分，取Top-K条款作为标准上下文注入抽取和生成流程。",
        "[0049] StandardClause对象至少包括standard_id、authority、version、clause_id、title_cn、title_en、scope、keywords、applicable_entity_types、required_fields、equivalent_refs、summary和effective_date字段。若检索到多个等效条款，系统保留authority和version差异，不把不同适航当局条款强行合并；若当前资料引用旧版本条款，系统记录version_mismatch事件，提示人工确认是否采用新版条款。",
        "[0050] 多角色协同流程包括知识提取器、本体推理器、安全评估器、追溯检查器、变更影响分析器和文档整理器。知识提取器负责候选实体和关系；本体推理器负责端点类型、层级归属和潜在链接；安全评估器负责危险、失效状态、安全目标、DAL建议和缓解措施；追溯检查器负责需求、设计、测试、维修记录和标准条款之间的覆盖关系；变更影响分析器沿图谱边输出上游、下游、横向、安全和培训影响；文档整理器只对已经通过校验的对象生成证据包和报告。",
        "[0051] 六层校验按固定顺序执行。第一层为格式校验，检查JSON字段、类型、必填项和枚举值；第二层为本体一致性校验，检查实体类型和关系端点是否匹配；第三层为标准引用校验，凡涉及安全目标、DAL、MEL、放行限制、维修记录和培训要求的对象必须绑定标准、规范或证据片段；第四层为证据跨度校验，检查source_span是否能回指到具体文档页码和段落；第五层为冲突检测，比较不同角色或不同文档对同一实体属性、原因关系或处置建议的差异；第六层为置信度阈值校验，低于阈值的对象进入人工复核队列。",
        "[0052] 冲突检测的执行规则为：若同一实体的entity_type不同，则标记type_conflict；若同一故障的原因方向相反，则标记causality_conflict；若同一维修动作对应不同放行限制，则标记release_conflict；若同一标准编号存在不同版本引用，则标记version_conflict。冲突对象不写入正式关系，只写入ConflictRecord，字段包括conflict_id、object_ids、conflict_type、evidence_ids、candidate_values、suggested_action和review_status。",
        "[0053] 经验闭环在校验后执行。系统将通过校验的故障分析、维修处置、标准引用、人工复核意见和审计记录组织为ExperienceEntry，字段包括experience_id、title、aircraft_type、task_type、applicable_conditions、content、evidence_ids、standard_refs、confidence、review_status、created_by和usage_count。经验条目先进入临时经验库，经工程人员审核后可提升为永久经验。后续遇到相同机型、相同系统、相似故障或相同标准条款时，系统把匹配经验作为上下文注入，但经验只作为辅助约束，不替代标准条款和原文证据。",
        "[0054] 本实施例的输出包括正式知识图谱、候选对象池、冲突记录、标准引用清单、经验条目、证据包和审计日志。若解析失败，系统输出parse_failed状态；若标准缺失，输出standard_gap状态；若模型不可用，系统切换至规则抽取和保守候选生成；若人工复核未完成，对象保持candidate状态，不进入正式故障树或最终答案。"
    ]),
    ("实施例3：围绕三份航空维修PDF的故障知识图谱、问答和复核输出", [
        "[0055] 本实施例以三份实际航空维修资料为输入，说明从PDF资料到故障知识图谱、故障树、因果树、跨文档证据包和维修问答的完整实施过程。输入资料包括《737典型故障说明R1（飞行版）.pdf》《M2-航空器维修（第一次修订）R1版2025年4月7日.pdf》和《航空器维修基础知识和实作培训规范-正式版.pdf》。三份资料分别提供故障案例、维修管理和培训实作规范，系统将其作为同一航空维修知识域的不同证据源处理。",
        "[0056] 对《737典型故障说明R1（飞行版）.pdf》，系统执行以下步骤：S301，解析故障目录和故障详情页，识别PACK OFF、BLEED OFF、DUAL BLEED OFF、发动机滑油、飞行操纵和起飞构型等故障主题；S302，为每个故障主题建立Episode对象；S303，在Episode下建立故障现象、运行条件、可能原因、排故动作、排除原因、维修建议和预防提示等Facet；S304，把每个Facet中的具体表述转化为FacetPoint并关联EntityRecord；S305，依据触发词和章节语义建立symptom_of、caused_by、requires_check、excluded_cause和maintenance_action等关系。",
        "[0057] Episode对象至少包括episode_id、aircraft_type、fault_topic、source_document、start_page、end_page、facets、primary_symptom、main_system和review_status字段。Facet对象至少包括facet_id、episode_id、facet_type、facet_points和evidence_ids字段，其中facet_type可取symptom、condition、possible_cause、excluded_cause、maintenance_action、standard_reference、training_tip等。FacetPoint对象至少包括point_id、facet_id、text、entity_ids、polarity、modality和source_span字段，modality用于区分confirmed、possible、candidate和excluded。",
        "[0058] 以PACK OFF为例，系统把“PACK OFF灯亮”“引气压力异常”“空调组件工作异常”“查阅MEL”“完成维修记录”分别抽取为故障现象、失效状态、部件状态、标准条款和维修动作实体。若同一段落同时出现“压力低”和“组件关断”，系统不直接将二者互为因果，而是检查上下文是否存在“导致、引起、由于、检查、排除”等触发词，并结合MEL或维修程序引用确定关系方向。最终输出的PACK OFF子图至少包含故障现象节点、引气系统节点、空调组件节点、MEL条目节点、维修动作节点和培训提示节点。",
        "[0059] 对BLEED OFF和DUAL BLEED OFF类故障，系统优先识别引气系统、发动机引气、APU引气、活门状态、驾驶舱指示、操作条件和放行限制。若资料中同时出现“现象”和“操作条件”，系统将操作条件作为condition节点，不将其误判为原因；若资料中出现“排除某原因”或“检查正常”，系统将该对象标记为excluded_cause，并在故障树中用排除标记展示，而不作为OR门下的基本事件。",
        "[0060] 对发动机滑油类故障，系统优先抽取滑油压力、滑油温度、滑油量、传感器、指示异常、检查动作和维修记录；对飞行操纵和起飞构型类故障，系统优先抽取构型状态、警告状态、飞控部件、操作程序、安全提示和培训要求。不同故障类型对应不同实体模板，但均使用同一EntityRecord和RelationRecord结构，保证后续检索、问答和审计统一。",
        "[0061] 对《M2-航空器维修（第一次修订）R1版2025年4月7日.pdf》，系统重点抽取维修作风、人为因素、维修记录、放行要求、工具管理、工卡执行和诚信管理等实体。维修文件中的record_requirement、human_factor和procedure_requirement与故障说明中的maintenance_action关联。例如，若PACK OFF处置需要记录维修动作，系统把该维修动作与M2文件中的维修记录要求建立requires_record关系；若BLEED OFF问题涉及交接和复核，系统把相关故障链路与人为因素、交接复核要求建立risk_control关系。",
        "[0062] 对《航空器维修基础知识和实作培训规范-正式版.pdf》，系统重点抽取培训科目、实作步骤、考核要点、安全注意事项和MEL查阅训练要求。培训规范中的training_item和assessment_requirement与故障案例中的故障主题和维修动作关联。例如，MEL查阅训练可与PACK OFF和BLEED OFF故障案例关联，工具使用和工卡执行训练可与维修动作节点关联，安全注意事项可与飞行操纵或起飞构型风险节点关联。",
        "[0063] 跨文档关联分为共指消解、标准桥接和图结构确认三步。共指消解统一“空调组件”“PACK组件”“组件关断”等近义表达；标准桥接把MEL条目、维修记录要求和培训规范条款连接到相同或相关的标准对象；图结构确认检查两个实体是否共享部件、系统、动作或标准引用。只有同时满足名称相似、类型兼容和证据支持的关联才写入正式图谱，其他关联保留为候选关系。",
        "[0064] 共指消解的判定条件为：normalized_name相似度超过第一阈值，entity_type相同或父类型兼容，至少共享一个系统、部件、标准编号、故障主题或维修动作。若名称相似但类型冲突，例如一个为部件、一个为培训科目，则不合并，只建立related_to候选关系。若两个实体来源于不同文档但共享同一MEL条目和同一故障主题，则提高same_as候选分数。",
        "[0065] 故障树生成以顶事件为入口。系统接收顶事件名称，例如“引气系统压力低”或“PACK OFF灯亮”，在图谱中定位对应failure_state节点；沿caused_by、failure_leads_to和requires_check关系向上展开原因节点；将可直接检查或可维修的部件、传感器、活门、管路和控制逻辑标记为基本事件；根据多个原因是否独立、共同作用或排除关系推断OR门、AND门或NOT标记。每个故障树节点均保留node_type、source_span、evidence_ids和review_status字段。",
        "[0066] 因果树生成以时间和证据强度为主线，从故障现象开始，沿confirmed优先、possible次之、candidate再次之的顺序组织因果链。若某一关系被737故障说明和M2维修文件同时支持，则提高其显示强度；若培训规范只提供操作提示而不支持因果判断，则作为training_tip挂接，不进入因果主链。由此可避免把培训要求误当作故障原因，也可避免把维修动作误当作故障现象。",
        "[0067] 用户问答时，系统首先创建QueryRecord，字段包括query_id、user_id、query_text、query_type、aircraft_type、priority、created_time、route、status和audit_id。若问题包含明确故障码、标准编号、部件名或MEL条目，则优先使用基线混合检索路径；若问题包含“为什么、如何排故、影响哪些、关联哪些标准”等多步推理表达，则使用Agentic多步路径；若问题要求综合三份资料或列出所有证据，则使用递归检索路径；若高级路径未启用或模型不可用，则降级到基线混合检索路径。",
        "[0068] 基线混合检索路径包括：对问题进行分词、同义词扩展和实体识别；同时检索BM25索引、向量索引和图谱邻接表；将候选证据按来源文档多样性、实体匹配度、标准引用完整性和页码有效性重排；生成EvidenceBundle。Agentic路径在每一步保存step_id、current_query、retrieved_evidence、decision、next_query和stop_reason；递归检索路径通过受控函数对证据分组、过滤和摘要。三种路径的最终答案均必须引用EvidenceBundle中的证据编号。",
        "[0069] 证据锚定生成器接收EvidenceBundle后，按固定模板输出答案：第一部分为直接结论或保守判断，第二部分为证据依据，第三部分为维修或培训提示，第四部分为需人工复核事项，第五部分为引用来源。对于PACK OFF问题，答案可列出故障现象、可能原因、检查动作、MEL或维修记录要求和培训提示；对于BLEED OFF问题，答案可强调引气系统状态、维修记录和人为因素；对于发动机滑油或飞行操纵问题，答案必须提示安全复核，不把模型输出作为放行依据。",
        "[0070] 图形界面复核包括文档列表、知识树、图谱视图、证据包视图、候选关系队列和审计日志视图。工程人员可查看节点来源页码、关系强度、标准引用和模型置信度；可对候选关系执行确认、驳回、合并、拆分或补充证据；可将复核结论写回经验库。系统对每次人工操作生成AuditRecord，包括操作者、时间、对象、操作前状态、操作后状态和理由。报告输出模块可生成故障分析报告、追溯矩阵、标准引用清单、维修培训提示和审计摘要，报告中的每一条结论均引用EvidenceRecord。",
        "[0071] 本实施例的输出对应说明书附图中的图7至图10，并可支撑图11和图12所示评测结果。若某一PDF页解析质量不足或来源证据冲突，系统在证据包中标注需要人工复核，而不是生成确定性维修结论。"
    ]),
    ("实施例4：国产计算卡适配和高吞吐token推理部署", [
        "[0072] 本实施例说明硬件适配层和高吞吐token推理链路的可实施方式。硬件适配层不限定单一厂商或单一开源项目，而是把CPU、GPU、NPU、DCU、MLU及其他异构设备抽象为统一DeviceProfile，并根据设备能力选择模型加载、算子后端、张量精度、内存池、并行方式和回退策略。可选设备包括但不限于昇腾NPU、海光DCU、寒武纪MLU、昆仑芯、摩尔线程计算卡以及通用GPU。",
        "[0073] 硬件部署步骤包括：S401，执行设备探测，读取设备数量、驱动版本、算子库版本、可用内存、通信拓扑和健康状态；S402，检查模型格式和后端兼容性，必要时执行格式转换、权重量化或张量并行切分；S403，为嵌入模型、重排序模型和生成模型分别选择后端；S404，创建请求队列、KV cache页池、prefix cache、输出流管理器和审计计数器；S405，发送健康检查请求，确认首token返回、流式输出和日志写入均可用。",
        "[0074] 模型加载策略由ModelPlan对象描述，字段包括model_id、model_role、backend、device_ids、precision_mode、quantization_mode、max_model_len、tensor_parallel_size、memory_budget、load_status和fallback_model。model_role可取embedding、rerank、generation或summary。若生成模型无法在目标计算卡上加载，系统可将generation角色回退到较小模型或CPU/GPU后端，同时保留embedding和rerank角色在加速设备上运行。",
        "[0075] 高吞吐推理请求被拆分为预填充阶段和逐token解码阶段。预填充阶段处理用户问题、检索证据、标准条款和历史经验上下文；逐token解码阶段连续生成答案。请求队列记录request_id、session_id、prompt_tokens、expected_output_tokens、evidence_size、priority、arrival_time、deadline和safety_level。调度器按输入长度、证据包大小、优先级、安全等级和设备剩余内存组成动态batch；每个解码步结束后，已完成请求释放缓存，新请求可加入批次，形成连续批处理。",
        "[0076] 动态batch的组织规则为：若多个请求的prompt_tokens差异小于第一阈值且safety_level相同，则优先合并；若请求共享同一故障主题、同一MEL条目或同一标准条款前缀，则优先合并以提高prefix cache命中率；若某请求涉及安全评估、DAL建议或放行限制，则降低其最大batch大小并启用更严格引用检查；若某请求超过deadline，则提升优先级或单独解码。调度器每轮输出active_batch、waiting_queue、finished_requests和evicted_cache_pages。",
        "[0077] 分页KV cache采用固定大小页块管理每个会话的键值缓存。CachePage对象包括page_id、request_id、layer_id、device_id、token_start、token_end、ref_count、last_access_time和reusable_flag。若多个请求共享相同标准条款、MEL条目、PACK OFF故障背景或培训规范前缀，系统通过prefix cache复用预填充结果，仅对新增问题和新增证据执行计算。缓存淘汰优先级为已结束会话、低优先级会话、可重建前缀、长时间未访问页和候选问答页。",
        "[0078] prefix cache的键值可由model_id、standard_refs、fault_topic、aircraft_type、evidence_hash和prompt_template_version组合生成。缓存命中后，系统仍检查证据版本和标准版本是否一致；若源PDF重新导入、标准条款版本变化或提示词模板变化，则缓存失效并重新预填充。该机制避免旧证据污染新答案。",
        "[0079] 量化和混合精度策略按任务风险和硬件能力选择。培训问答、一般标准检索和资料导航可采用较低比特量化或混合精度；涉及安全评估、DAL建议、放行限制或复杂故障诊断的问题可采用较高精度、较小batch和更严格证据引用。若设备内存不足，系统依次尝试降低batch大小、缩短最大输出长度、启用量化模型、关闭低优先级会话、切换到CPU/GPU或低并发模式，并记录降级原因。",
        "[0080] 面向三份航空维修PDF的并发问答过程如下：第一用户提出PACK OFF灯亮排故问题，系统检索737故障说明中的故障条目、M2文件中的维修记录和培训规范中的MEL训练要求；第二用户提出BLEED OFF与人为因素问题，系统检索引气系统、维修作风和交接复核证据；第三用户提出发动机滑油或飞行操纵风险问题，系统检索故障说明、安全提示和培训规范。调度器把三类请求按证据长度和生成长度组成动态batch，在国产计算卡上执行重排序、摘要和答案生成，并向各用户流式返回。",
        "[0081] 审计统计包括model_version、hardware_backend、device_id、precision_mode、quantization_mode、batch_size、input_tokens、output_tokens、first_token_latency、total_latency、tokens_per_second、cache_hit_rate、kv_pages_used、fallback_event和evidence_ids。上述统计与用户、问题、证据包和答案摘要绑定，形成可追溯日志。对于安全相关问题，系统在答案中保留人工复核提示，并将审计日志用于后续性能调优、硬件容量规划和安全复盘。",
        "[0082] 异常回退分为四级。一级回退为同设备降级，例如降低batch大小或最大输出长度；二级回退为同类设备切换，例如从一张计算卡切换到另一张可用计算卡；三级回退为异构设备切换，例如从NPU/DCU/MLU切换到GPU或CPU；四级回退为保守模板模式，即不调用生成模型，只根据检索证据输出引用清单和人工复核提示。每次回退均写入AuditRecord，并在用户界面显示当前回答的生成模式。",
        "[0083] 本实施例的输出为可运行的国产化高吞吐维修问答服务。该服务可在内网或离线环境中部署；当国产计算卡可用时，生成、重排序和证据摘要可由计算卡加速；当计算卡不可用时，检索、证据包整理和保守模板回答仍可在CPU/GPU或低并发模式下运行。由此，本发明在不同单位软硬件条件下均可保持基本可用，并在有加速设备时提升多用户并发场景下的token吞吐。"
    ]),
]


def find_para(doc: Document, prefix: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i
    raise RuntimeError(f"marker not found: {prefix}")


def remove_para(p):
    p._element.getparent().remove(p._element)


def add_after(anchor, text: str, bold: bool = False):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    p = anchor._parent.add_paragraph()
    p._p = new_p
    r = p.add_run(text)
    r.font.name = "宋体"
    r.font.size = Pt(10.5)
    r.bold = bold
    return p


def rewrite():
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)
    start = find_para(doc, "第四部分：具体实施方式")
    end = find_para(doc, "第五部分：发明效果")
    for p in list(doc.paragraphs[start + 1:end]):
        remove_para(p)
    anchor = doc.paragraphs[start]
    flat = []
    for title, paras in IMPLEMENTATION:
        if title == "导言":
            flat.extend((p, False) for p in paras)
        else:
            flat.append((title, True))
            flat.extend((p, False) for p in paras)
    for content, bold in reversed(flat):
        add_after(anchor, content, bold=bold)
    doc.save(OUT)


def update_report():
    note = f"""

## 第四部分详化版

补充时间：{datetime.now():%Y-%m-%d %H:%M:%S}

在保持4个大实施例结构不变的前提下，已进一步扩写第四部分，新增更细的部署准备、配置字段、硬件探测判定、数据对象字段、PDF解析质量判断、chunk类型规则、实体归一、关系强度、标准条款对象、冲突类型、三份PDF故障抽取流程、故障树/因果树生成规则、问答路由、证据包模板、GUI复核、动态batch、分页KV cache、prefix cache、量化策略和四级异常回退。详化结果输出为：{OUT}
"""
    REPORT.write_text(REPORT.read_text(encoding="utf-8") + note, encoding="utf-8")


def checks():
    doc = Document(OUT)
    txt = "\n".join(p.text for p in doc.paragraphs)
    bad = ["综合实施例", "软件设计说明书", "第九部分", "第十部分", "待补充", "TODO", "占位图", "示意图占位"]
    with zipfile.ZipFile(OUT) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    return {
        "bad_terms": [b for b in bad if b in txt],
        "embodiments": sum(1 for p in doc.paragraphs if p.text.strip().startswith("实施例")),
        "paragraphs": len(doc.paragraphs),
        "media": len(media),
        "out": str(OUT),
    }


if __name__ == "__main__":
    rewrite()
    update_report()
    print(checks())
