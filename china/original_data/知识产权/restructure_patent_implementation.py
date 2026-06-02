from __future__ import annotations

import copy
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement


IP = Path(__file__).resolve().parent
DOCX = IP / "发明专利_重构申请文档_知产权综合润色版_附图硬件增强版.docx"
BACKUP = IP / f"发明专利_重构申请文档_知产权综合润色版_附图硬件增强版_结构调整前备份_{datetime.now():%Y%m%d_%H%M%S}.docx"
REPORT = IP / "专利文档_附图硬件增强_修改报告.md"


def text(p):
    return p.text.strip()


def find_para(doc, prefix):
    for i, p in enumerate(doc.paragraphs):
        if text(p).startswith(prefix):
            return i
    raise RuntimeError(f"marker not found: {prefix}")


def insert_paragraph_after(paragraph, content):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = paragraph._parent.add_paragraph()
    p._p = new_p
    r = p.add_run(content)
    r.font.name = "宋体"
    r.font.size = Pt(10.5)
    return p


def remove_paragraph_element(p_el):
    parent = p_el.getparent()
    if parent is not None:
        parent.remove(p_el)


def rebuild_implementation_section(doc):
    # Move all useful implementation material into the fourth section.
    fifth_idx = find_para(doc, "第五部分：发明效果")
    section8_idx = find_para(doc, "第八部分：综合实施例补充")
    section9_idx = find_para(doc, "第九部分：软件设计说明书支撑摘录")
    section10_idx = find_para(doc, "第十部分：专利交底书补充材料归并")
    first_comp_after_10 = find_para(doc, "综合实施例10")
    docs_section_idx = find_para(doc, "第十一部分：围绕docs目录PDF资料的具体实施方式")
    project_section_idx = find_para(doc, "第十一部分：项目解析后的详细实施方式")
    support_idx = find_para(doc, "实施例56：发明内容与实施例支撑关系")
    section12_idx = find_para(doc, "第十二部分：发明内容支撑关系表")

    useful_ranges = [
        (section8_idx + 1, section9_idx),
        (first_comp_after_10, docs_section_idx),
        (docs_section_idx + 1, project_section_idx),
        (project_section_idx + 1, support_idx),
    ]
    useful_paras = []
    for start, end in useful_ranges:
        useful_paras.extend(p._p for p in doc.paragraphs[start:end])

    # Capture the original fragmented tail before inserting cloned paragraphs; paragraph
    # indexes change as soon as we insert before the fifth section.
    to_remove = [p._p for p in doc.paragraphs[section8_idx:section12_idx]]

    fifth_el = doc.paragraphs[fifth_idx]._p
    body = doc.element.body
    for p_el in useful_paras:
        body.insert(body.index(fifth_el), copy.deepcopy(p_el))

    # Remove old fragmented sections: section 8, software design excerpt, patent disclosure material,
    # duplicate section 11 headings, support-relation meta section and their original paragraphs.
    for p_el in to_remove:
        remove_paragraph_element(p_el)


def renumber_and_clean(doc):
    fourth = find_para(doc, "第四部分：具体实施方式")
    fifth = find_para(doc, "第五部分：发明效果")
    counter = 0
    for p in doc.paragraphs[fourth + 1:fifth]:
        t = text(p)
        m = re.match(r"^(综合实施例|实施例)(\d+续?|\d+)?：(.+)$", t)
        if m:
            counter += 1
            p.text = f"实施例{counter}：{m.group(3)}"
            for r in p.runs:
                r.bold = True
        elif t.startswith("以下结合附图和实施例"):
            p.text = "以下结合附图和连续实施例，对本发明作进一步详细说明。下列实施例均围绕同一技术方案展开，后续实施例是对前述方法步骤、系统模块、数据对象、航空维修PDF处理流程、硬件部署和审计回退机制的进一步细化，不构成彼此割裂的独立方案。"

    # Rename remaining later sections so the final document no longer has a separate eighth/ninth
    # implementation supplement or software-design excerpt.
    for p in doc.paragraphs:
        t = text(p)
        if t.startswith("第六部分：附图说明"):
            p.text = "第六部分：附图说明"
        elif t.startswith("第七部分：摘要"):
            p.text = "第七部分：摘要"
        elif t.startswith("第十三部分：说明书附图"):
            p.text = "第八部分：说明书附图"


DETAILS = {
    "系统整体架构与部署环境": [
        "在部署时，系统先建立资料输入区、解析缓存区、知识图谱区、标准条款区、经验库区和审计日志区。资料输入区只保存原始文件及文件指纹，解析缓存区保存按页码、章节、表格和段落拆分后的中间结果，知识图谱区保存实体、关系和关联强度，标准条款区保存可检索的条款对象，经验库区保存经审核的经验条目，审计日志区保存每一次模型调用、检索调用和人工复核操作。",
        "当用户上传航空维修PDF时，系统为每个文件生成document_id，并将页码、章节标题、段落序号、字符偏移、表格坐标和图片说明统一写入source_span。后续实体抽取、关系抽取、证据生成和人工复核均引用该source_span，从而保证任何结论均能回指到原文位置。"
    ],
    "B737典型故障说明文档的端到端知识提取": [
        "以PACK OFF故障为例，系统首先识别故障现象类片段，如驾驶舱告警、组件状态、引气压力和空调组件状态；其次识别处置步骤类片段，如检查引气源、检查组件活门、查阅MEL和记录维修措施；再次识别限制条件类片段，如放行限制、重复故障条件和人工复核要求。上述片段分别映射为failure_state、component、maintenance_action、standard_clause和training_item等实体。",
        "关系构建时，系统不把同一段文字中的所有实体两两相连，而是根据句法触发词、章节位置、标准条款引用和M-flow因果边界判断关系方向。例如“压力低导致PACK OFF灯亮”被组织为failure_leads_to关系，“查阅MEL 21-52-01”被组织为refers_to_standard关系，“完成维修记录”被组织为requires_record关系。"
    ],
    "跨文档关联与标准条款映射": [
        "跨文档关联分三轮执行。第一轮执行名称归一，将PACK OFF、空调组件关断、组件失效等不同表述归并到同一候选实体；第二轮执行标准桥接，将故障说明中的MEL条目、M2教材中的维修记录要求和培训规范中的实作训练要求映射到统一标准或规范对象；第三轮执行图结构确认，当两个候选实体共享部件、故障模式、维修动作或培训科目时，提高其关联强度并写入跨文档边。",
        "对于关联证据不足的候选边，系统保留candidate状态和证据来源，但不进入正式故障树或因果树。工程师可在GUI中查看候选边的来源页码、触发词、标准条款和模型置信度，确认后再提升为正式关系。"
    ],
    "面向国产计算卡的航空维修问答高吞吐推理部署方法": [
        "请求调度器将每个维修问答请求拆分为检索阶段、证据整理阶段、生成预填充阶段和逐token解码阶段。检索阶段可并行访问BM25索引、向量索引和图谱邻接表；证据整理阶段将来源页码、标准条款、故障实体和培训提示压缩为证据包；预填充阶段复用热点前缀缓存；解码阶段以连续批处理方式合并多个会话的下一token计算。",
        "分页KV cache采用会话标识、页号、层号和设备号四级索引。每个会话的缓存页记录引用计数、最近访问时间、所属证据包和是否可复用。若设备内存不足，系统优先回收已结束会话、低优先级培训问答会话和可由prefix cache重建的缓存页；若仍不足，则降低批次大小或切换到短上下文模式。",
        "审计日志除记录最终答案外，还记录硬件后端、模型版本、量化方式、batch大小、缓存命中率、首token延迟、总token数和降级事件。对于PACK OFF、BLEED OFF、发动机滑油和飞行操纵等安全相关问题，系统将答案标记为“辅助建议”，并输出人工复核项，避免把模型生成内容直接作为维修放行依据。"
    ],
}


def add_detail_paragraphs(doc):
    fourth = find_para(doc, "第四部分：具体实施方式")
    fifth = find_para(doc, "第五部分：发明效果")
    # Iterate backwards so insertion does not disturb later anchors.
    for idx in range(fifth - 1, fourth, -1):
        p = doc.paragraphs[idx]
        t = text(p)
        for key, paras in DETAILS.items():
            if t.startswith("实施例") and key in t:
                anchor = p
                for detail in reversed(paras):
                    insert_paragraph_after(anchor, detail)
                break


REMOVE_TITLE_KEYWORDS = [
    "创造性三步法论证",
    "软件著作权",
    "面向专利申请",
    "与现有技术",
    "源程序材料转化",
    "申请文本",
    "英文发布材料",
    "知识产权目录材料",
    "不纳入正式保护范围",
    "面向代理人",
    "综合润色",
    "申请提交前",
]


def remove_meta_embodiments(doc):
    fourth = find_para(doc, "第四部分：具体实施方式")
    fifth = find_para(doc, "第五部分：发明效果")
    paras = doc.paragraphs
    remove = []
    i = fourth + 1
    while i < fifth:
        t = text(paras[i])
        if t.startswith("实施例") and any(k in t for k in REMOVE_TITLE_KEYWORDS):
            j = i + 1
            while j < fifth and not text(paras[j]).startswith("实施例"):
                j += 1
            remove.extend(p._p for p in paras[i:j])
            i = j
        else:
            i += 1
    for p_el in remove:
        remove_paragraph_element(p_el)


def final_renumber(doc):
    fourth = find_para(doc, "第四部分：具体实施方式")
    fifth = find_para(doc, "第五部分：发明效果")
    n = 0
    for p in doc.paragraphs[fourth + 1:fifth]:
        t = text(p)
        if t.startswith("实施例") and "：" in t:
            n += 1
            p.text = re.sub(r"^实施例\d+：", f"实施例{n}：", t)
            for r in p.runs:
                r.bold = True


def update_report():
    if not REPORT.exists():
        return
    extra = f"""

## 结构调整补充

补充时间：{datetime.now():%Y-%m-%d %H:%M:%S}

根据复核意见，已将原“第八部分：综合实施例补充”全部并入“第四部分：具体实施方式”，并删除“第九部分：软件设计说明书支撑摘录”和“第十部分：专利交底书补充材料归并”等材料汇编式内容。原综合实施例已统一改名并连续编号为“实施例”，避免说明书正文出现软件著作权说明书、交底书或代理布局建议等非专利正文痕迹。另对系统部署、B737故障抽取、跨文档关联和国产计算卡高吞吐推理实施例补充了更细的步骤、数据对象、调度和审计描述。
"""
    REPORT.write_text(REPORT.read_text(encoding="utf-8") + extra, encoding="utf-8")


def checks(docx_path):
    doc = Document(docx_path)
    txt = "\n".join(p.text for p in doc.paragraphs)
    bad = [
        "第八部分：综合实施例补充",
        "第九部分：软件设计说明书支撑摘录",
        "第十部分：专利交底书补充材料归并",
        "软件设计说明书",
        "综合实施例",
    ]
    with zipfile.ZipFile(docx_path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    return {
        "bad_terms": [b for b in bad if b in txt],
        "embodiments": sum(1 for p in doc.paragraphs if p.text.strip().startswith("实施例")),
        "media": len(media),
        "paragraphs": len(doc.paragraphs),
    }


def main():
    shutil.copy2(DOCX, BACKUP)
    doc = Document(DOCX)
    rebuild_implementation_section(doc)
    renumber_and_clean(doc)
    add_detail_paragraphs(doc)
    remove_meta_embodiments(doc)
    final_renumber(doc)
    doc.save(DOCX)
    update_report()
    print(f"backup {BACKUP}")
    print(checks(DOCX))


if __name__ == "__main__":
    main()
