from __future__ import annotations

import json
import zipfile
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = PROJECT_ROOT / "deliverables" / "AgriKB_DualTrack_FirstPrize_Materials_20260523"
TODAY = date(2026, 5, 23)

ACCENT = "2F7D52"
DARK = "17211D"
MUTED = "5F6F63"
LIGHT = "EEF5E9"
SECONDARY = "B6852D"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    paragraph.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color, before, after in [
        ("Heading 1", 18, ACCENT, 18, 8),
        ("Heading 2", 13.5, DARK, 12, 6),
        ("Heading 3", 11.5, SECONDARY, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    core = doc.core_properties
    core.title = title
    core.subject = "AgriKB 数智新农人经营中枢交付材料"
    core.author = "AgriKB Project"
    core.comments = "Generated locally for the 2026 Jurong young talent competition delivery package."


def add_cover(doc: Document, title: str, subtitle: str, version: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("AgriKB")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(DARK)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    p = doc.add_paragraph()
    run = p.add_run(subtitle)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    table = doc.add_table(rows=4, cols=2)
    table.autofit = False
    labels = [
        ("产品定位", "数智信息技术 × 新农人融合项目"),
        ("应用场景", "江苏句容县域农业治理、台湾农业开放数据样板、农户经营决策"),
        ("材料版本", version),
        ("生成日期", TODAY.isoformat()),
    ]
    for row, (label, value) in zip(table.rows, labels):
        set_cell_text(row.cells[0], label, True)
        set_cell_text(row.cells[1], value)
        set_cell_shading(row.cells[0], LIGHT)
    doc.add_paragraph()
    add_callout(
        doc,
        "交付说明",
        "本文档为本地项目交付稿，可直接用于演示、申报和二次完善。专利与软著内容属于初稿，需要在正式提交前由知识产权代理人或法务进行权利要求、署名、日期和材料一致性复核。",
    )
    doc.add_section(WD_SECTION_START.NEW_PAGE)


def add_callout(doc: Document, label: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(label + "：")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run = paragraph.add_run(body)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10)
    doc.add_paragraph()


def add_bullets(doc: Document, items: Iterable[str], style: str = "List Bullet") -> None:
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, True)
        set_cell_shading(cell, LIGHT)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            set_cell_text(row.cells[idx], value)
    doc.add_paragraph()


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = "AgriKB 数智新农人经营中枢 | 交付材料 | " + TODAY.isoformat()


def save_doc(doc: Document, path: Path) -> None:
    add_footer(doc)
    doc.save(path)


def write_markdown(path: Path, title: str, sections: list[tuple[str, list[str]]]) -> None:
    lines = [f"# {title}", "", f"- 版本：V1.0 交付稿", f"- 日期：{TODAY.isoformat()}", ""]
    for heading, paragraphs in sections:
        lines.extend([f"## {heading}", ""])
        for paragraph in paragraphs:
            lines.append(paragraph)
            lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def build_user_manual() -> Path:
    title = "AgriKB 数智新农人经营中枢使用说明书"
    doc = Document()
    configure_document(doc, title)
    add_cover(doc, title, "面向新农人、农业局、农业企业的一线操作手册", "V1.0 交付稿")

    doc.add_heading("1. 产品入口", level=1)
    add_bullets(
        doc,
        [
            "在项目根目录双击 RUN_AGRIKB.bat 启动；默认访问地址为 http://127.0.0.1:8010/ui/。",
            "如果农业局或合作社内网部署，手机和其他电脑可访问服务器局域网 IP 加端口，例如 http://192.168.x.x:8010/ui/。",
            "系统启动后，页面右上或标题区会显示健康状态、当前模型、知识库条目数和是否启用真实模型回答。",
        ],
    )

    doc.add_heading("2. 标准问答流程", level=1)
    add_numbered(
        doc,
        [
            "在输入框描述作物、地点、经营目标和约束，例如“句容草莓这周是否适合采收，如何匹配电商和政策补贴”。",
            "按需打开“联网搜索”，系统会把实时天气、政策、新闻和市场信息作为增量情报纳入回答。",
            "提交后等待模型生成：回答会优先给经营判断、政策匹配、市场/收购建议，再给可执行步骤。",
            "默认隐藏证据链和知识树；需要审计、申报或复核时，点击展开即可查看引用来源、知识图谱节点和推理依据。",
        ],
    )

    doc.add_heading("3. 三类用户怎么用", level=1)
    add_table(
        doc,
        ["用户", "典型问题", "输出重点"],
        [
            ["新农人/合作社", "采收、病虫害、销售渠道、补贴匹配", "今天做什么、卖给谁、价格/渠道风险、需准备的材料"],
            ["农业局", "产业扶持、培训、项目筛选、台账治理", "政策条款匹配、主体清单、风险提示、可追溯证据"],
            ["农企/采购商", "产地采购、品质认证、电商入驻、供应链", "供需判断、认证/溯源、采购窗口、合作路径"],
        ],
    )

    doc.add_heading("4. 数据与资料补充", level=1)
    add_bullets(
        doc,
        [
            "可拖拽 PDF、DOCX、TXT、Markdown、CSV、XLSX、JSON、XML、ZIP 等资料进入“补充农业文档”区域。",
            "系统会自动抽取、分段、入库，并将新资料纳入农业知识图谱；适合导入农业局政策文件、合作社台账、收购表、认证清单。",
            "当前知识库已纳入 AGROVOC、Crop Ontology、Google Earth Engine agriculture 标签资料、台湾农业开放数据样板、用户提供 COA_OpenData 数据，以及句容创业大赛政策要点。",
        ],
    )

    doc.add_heading("5. 演示时建议提问", level=1)
    add_bullets(
        doc,
        [
            "“我在句容做草莓种植，今天要不要采收，怎么结合天气、收购、电商和新农人政策安排？”",
            "“农业局要筛选一批新农人项目，如何按数智信息技术和新农人赛道做评分？”",
            "“台湾有机农业经营者数据能给句容农产品品牌和认证体系带来什么启发？”",
            "“请生成一份面向创业大赛答辩的技术壁垒、商业模式和落地计划。”",
        ],
    )

    doc.add_heading("6. 故障处理", level=1)
    add_table(
        doc,
        ["现象", "处理方式"],
        [
            ["模型回答很慢", "确认 Ollama 或 Kimi 后台可用；关闭不必要的联网搜索；减少一次提问中的长文件数量。"],
            ["Kimi 调用失败", "检查 .env 中的 Moonshot/Kimi 地址、模型名和密钥；系统会自动回退到本地模型。"],
            ["手机无法访问", "确认电脑和手机在同一网络；Windows 防火墙放行 Python/端口 8010；用电脑局域网 IP 访问。"],
            ["证据链为空", "尝试打开联网搜索或补充本地政策/数据文件；问题越具体，检索命中越稳定。"],
        ],
    )

    path = OUT_DIR / "01_AgriKB_使用说明书.docx"
    save_doc(doc, path)
    write_markdown(
        OUT_DIR / "01_AgriKB_使用说明书.md",
        title,
        [
            ("启动入口", ["双击 RUN_AGRIKB.bat，访问 http://127.0.0.1:8010/ui/。"]),
            ("问答流程", ["输入农业经营问题，按需打开联网搜索，默认查看结论，可展开证据链和知识树。"]),
            ("适用用户", ["新农人、农业局、合作社、农企、电商助农服务商。"]),
        ],
    )
    return path


def build_product_description() -> Path:
    title = "AgriKB 数智新农人经营中枢产品说明书"
    doc = Document()
    configure_document(doc, title)
    add_cover(doc, title, "比赛申报、商务沟通、政府汇报通用版", "V1.0 交付稿")

    doc.add_heading("1. 产品定位", level=1)
    add_callout(
        doc,
        "一句话",
        "AgriKB 是面向县域农业的 AI 经营智能体，用数智信息技术把政策、市场、气象、地理环境、作物知识和本地农业数据整合成可执行的经营决策。",
    )
    add_bullets(
        doc,
        [
            "赛道定位：数智信息技术 × 新农人双赛道融合，既能展示 AI/RAG/知识图谱技术壁垒，又能落到农业经营。",
            "落地定位：以江苏句容为示范场景，以台湾农业开放数据作为标准化样板，服务县域农业治理和新农人创业。",
            "交付定位：不是网页演示，而是可本地运行、可联网检索、可补充资料、可审计证据链的完整软件。",
        ],
    )

    doc.add_heading("2. 目标用户与核心价值", level=1)
    add_table(
        doc,
        ["对象", "当前痛点", "AgriKB 价值"],
        [
            ["新农人", "懂种植但不一定懂政策、市场、电商和数据工具", "把“今天怎么做”转成采收、销售、补贴、风险的行动清单"],
            ["农业局", "政策、主体、项目、台账分散，难以快速形成产业判断", "提供项目筛选、政策匹配、产业研判和证据化报告"],
            ["农企", "采购、认证、溯源、渠道信息碎片化", "辅助做产地合作、收购窗口、品质背书和供应链判断"],
        ],
    )

    doc.add_heading("3. 核心功能", level=1)
    add_bullets(
        doc,
        [
            "农业知识库问答：围绕作物、病虫害、性状、认证、主体、政策、市场和渠道进行模型分析。",
            "实时情报增强：接入联网搜索，把新闻政策、气象、产业扶持、市场和收购信息纳入回答。",
            "证据链与知识树：默认隐藏，必要时展开，满足农业局审计、比赛答辩和企业尽调需要。",
            "本地资料入库：支持导入农业局文件、Excel 台账、CSV 开放数据、PDF 政策文件等。",
            "移动端适配：支持手机浏览器/PWA 安装，适合田间、合作社、培训会现场使用。",
        ],
    )

    doc.add_heading("4. 技术架构", level=1)
    add_table(
        doc,
        ["层级", "能力", "说明"],
        [
            ["数据层", "开放数据 + 本地资料 + 实时搜索", "AGROVOC、Crop Ontology、GEE、台湾开放数据、句容政策、用户文件。"],
            ["知识层", "分段、索引、知识图谱", "将作物、主体、政策、市场、证据来源组织为可检索知识空间。"],
            ["模型层", "Kimi/Moonshot 兼容接口 + 本地 Ollama 回退", "保证能真正调用模型生成综合分析，避免固定话术。"],
            ["应用层", "问答、证据链、知识树、报告、朗读", "面向农民、农业局、企业的不同工作流。"],
        ],
    )

    doc.add_heading("5. 商业模式", level=1)
    add_bullets(
        doc,
        [
            "政府端：县域农业知识中台、项目筛选系统、政策兑现辅助、培训与运维服务。",
            "企业端：农业采购与产地协同、认证溯源、供应链情报、品牌/电商运营助手。",
            "新农人端：轻量订阅、合作社集体采购、培训营工具包、助农服务渠道接入。",
            "数据服务：区域作物经营模型、政策知识包、市场收购情报包、可追溯报告生成。",
        ],
    )

    doc.add_heading("6. 竞赛第一名叙事", level=1)
    add_bullets(
        doc,
        [
            "不是单点 AI 聊天，而是县域农业“数据、模型、知识、经营”的闭环。",
            "技术上有 RAG/GraphRAG、知识图谱、实时搜索、多模型后台和本地部署能力。",
            "产业上直击新农人的采收、销售、政策、渠道、融资和培训问题。",
            "政府侧可做产业扶持和项目筛选，商业侧可接采购、电商、认证和品牌运营。",
            "演示可现场提问、现场引用政策与市场信息、现场展开证据链，可信度强。",
        ],
    )

    path = OUT_DIR / "02_AgriKB_产品说明书.docx"
    save_doc(doc, path)
    write_markdown(
        OUT_DIR / "02_AgriKB_产品说明书.md",
        title,
        [
            ("定位", ["县域农业 AI 经营智能体，数智信息技术 × 新农人融合项目。"]),
            ("价值", ["把政策、市场、气象、地理环境和开放农业数据转成经营判断。"]),
            ("优势", ["可本地运行、可联网检索、可导入资料、可展开证据链和知识树。"]),
        ],
    )
    return path


def build_patent_draft() -> Path:
    title = "AgriKB 专利交底书及申请文件初稿"
    doc = Document()
    configure_document(doc, title)
    add_cover(doc, title, "一种面向县域农业经营的多源数据检索增强生成与知识图谱决策系统及方法", "专利初稿 V0.1")

    doc.add_heading("1. 建议专利名称", level=1)
    add_bullets(
        doc,
        [
            "一种面向县域农业经营的多源数据检索增强生成与知识图谱决策系统及方法",
            "一种数智新农人经营智能体系统、方法、电子设备及存储介质",
        ],
    )

    doc.add_heading("2. 技术领域", level=1)
    doc.add_paragraph(
        "本发明涉及人工智能、农业信息化、知识图谱、检索增强生成、县域产业治理和农业经营决策技术领域，尤其涉及一种面向新农人、农业主管部门和农产品经营主体的多源农业数据融合、证据链生成和经营决策方法。"
    )

    doc.add_heading("3. 背景技术", level=1)
    add_bullets(
        doc,
        [
            "农业经营决策同时依赖作物知识、气象、地理环境、政策、市场、认证、主体台账和渠道信息，但现有工具通常只覆盖其中一类数据。",
            "通用大模型可回答农业问题，但缺乏本地政策和经营主体数据，容易出现不可追溯、不可审计、难以落地的问题。",
            "农业局和新农人在项目申报、产业扶持、采收销售、电商协同等场景中，需要同时获得结论、证据、风险和行动建议。",
        ],
    )

    doc.add_heading("4. 发明内容", level=1)
    add_callout(
        doc,
        "核心思路",
        "系统将农业开放知识、本地农业资料、实时搜索情报和用户问题统一映射到农业知识图谱与证据对象中，再通过多路检索、重排序和大模型生成，输出包含经营判断、政策匹配、市场建议和可展开证据链的答案。",
    )
    add_table(
        doc,
        ["模块", "功能"],
        [
            ["多源数据采集模块", "采集 AGROVOC、Crop Ontology、GEE、开放数据、政策 PDF、CSV 台账和网页搜索结果。"],
            ["农业语义规范模块", "将作物、性状、认证、主体、区域、政策、市场、渠道统一为农业语义实体。"],
            ["证据对象构建模块", "为每个文本块、表格行、搜索摘要生成来源、时间、主题和可信度标签。"],
            ["混合检索与图谱扩展模块", "结合关键词、向量、图谱邻接关系和任务意图进行候选证据召回。"],
            ["经营决策生成模块", "调用大模型生成结论，并强制覆盖政策、市场、收购、风险和行动建议。"],
            ["审计展示模块", "默认隐藏证据链和知识树，用户可按需展开查看来源和推理路径。"],
        ],
    )

    doc.add_heading("5. 方法步骤", level=1)
    add_numbered(
        doc,
        [
            "接收用户输入的农业经营问题，并识别地点、作物、主体、时间、经营目标和风险约束。",
            "根据问题意图触发本地知识库检索、图谱邻域扩展、实时网页搜索和气象/环境信息查询。",
            "将检索结果标准化为证据对象，记录来源、发布时间、地域、适用主体、数据类型和置信标签。",
            "对证据对象进行去重、重排序和冲突检测，形成政策证据、市场证据、作物证据、气象环境证据和经营案例证据集合。",
            "构造大模型提示词，要求模型输出经营判断、政策匹配、市场/收购信息、执行步骤和风险提示。",
            "将模型答案与证据链、知识图谱节点、用户会话记忆和可下载报告关联存储。",
        ],
    )

    doc.add_heading("6. 有益效果", level=1)
    add_bullets(
        doc,
        [
            "降低新农人获取政策和市场信息的门槛。",
            "提高农业局项目筛选、政策扶持和产业研判的证据化程度。",
            "通过默认隐藏、按需展开的证据链机制兼顾易用性与审计性。",
            "支持本地模型回退和私有化部署，适配县域农业数据安全要求。",
            "可面向不同地区替换本地数据包，具备复制推广能力。",
        ],
    )

    doc.add_heading("7. 权利要求书初稿", level=1)
    claims = [
        "一种面向县域农业经营的多源数据检索增强生成与知识图谱决策系统，包括多源数据采集模块、农业语义规范模块、证据对象构建模块、混合检索与图谱扩展模块、经营决策生成模块以及审计展示模块。",
        "根据权利要求1所述的系统，其中多源数据采集模块用于采集农业控制词表、作物本体、遥感数据目录、政府开放数据、本地政策文件、农业主体台账以及实时网页搜索结果。",
        "根据权利要求1所述的系统，其中农业语义规范模块将作物、性状、区域、主体、政策、认证、渠道、价格和风险映射为统一农业语义实体。",
        "根据权利要求1所述的系统，其中证据对象构建模块为文本块、表格行、搜索摘要和文件片段生成来源、时间、地域、主题、可信度和适用主体标签。",
        "根据权利要求1所述的系统，其中混合检索与图谱扩展模块基于关键词检索、向量检索、图谱邻域和用户意图路由进行候选证据召回。",
        "根据权利要求1所述的系统，其中经营决策生成模块调用大语言模型，并在提示词中约束答案至少包含经营判断、政策匹配、市场或收购建议、风险提示和执行步骤。",
        "根据权利要求1所述的系统，其中审计展示模块默认隐藏证据链和知识树，并响应用户操作展示证据来源、图谱节点和推理路径。",
        "一种面向县域农业经营的多源数据检索增强生成方法，包括接收问题、意图识别、多源检索、证据对象构建、证据重排序、模型生成、答案审计展示和会话记忆更新步骤。",
        "一种电子设备，包括处理器和存储器，所述存储器存储程序，所述程序被处理器执行时实现权利要求8所述的方法。",
        "一种计算机可读存储介质，其上存储计算机程序，所述程序被处理器执行时实现权利要求8所述的方法。",
    ]
    add_numbered(doc, claims)

    doc.add_heading("8. 摘要初稿", level=1)
    doc.add_paragraph(
        "本发明公开了一种面向县域农业经营的多源数据检索增强生成与知识图谱决策系统及方法。该系统采集农业开放知识、本地农业资料、政府政策文件、农业主体台账和实时网页搜索结果，将其标准化为农业语义实体和证据对象，通过关键词、向量、图谱邻域和意图路由进行混合检索，再调用大语言模型生成包含经营判断、政策匹配、市场收购建议、风险提示和执行步骤的答案。系统提供默认隐藏、按需展开的证据链和知识树展示机制，兼顾新农人易用性、农业主管部门审计性和农企经营决策需求，适用于县域农业治理、农产品经营、电商助农和新农人创业服务。"
    )

    path = OUT_DIR / "03_AgriKB_专利交底书_初稿.docx"
    save_doc(doc, path)
    write_markdown(
        OUT_DIR / "03_AgriKB_专利交底书_初稿.md",
        title,
        [
            ("建议名称", ["一种面向县域农业经营的多源数据检索增强生成与知识图谱决策系统及方法。"]),
            ("核心创新", ["多源农业证据对象、农业知识图谱扩展、面向经营判断的 RAG 生成、可展开证据链。"]),
            ("注意", ["正式提交前需由代理师检索现有技术并重写权利要求。"]),
        ],
    )
    return path


def build_software_copyright() -> Path:
    title = "AgriKB 软件著作权登记材料初稿"
    doc = Document()
    configure_document(doc, title)
    add_cover(doc, title, "软件著作权登记说明、功能说明和材料清单", "软著初稿 V0.1")

    doc.add_heading("1. 软件基本信息", level=1)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["软件全称", "AgriKB 数智新农人经营中枢软件"],
            ["软件简称", "AgriKB"],
            ["版本号", "V1.0"],
            ["软件类型", "农业知识库、智能问答、政策市场情报与经营决策支持软件"],
            ["运行环境", "Windows 10/11、Python 3.11、浏览器；移动端通过同网访问或 PWA 使用"],
            ["开发语言", "Python、JavaScript、HTML、CSS、SQLite、YAML/JSON 配置"],
        ],
    )

    doc.add_heading("2. 软件功能说明", level=1)
    add_bullets(
        doc,
        [
            "农业知识库构建：导入农业开放数据、本地政策文件、CSV/Excel 台账和网页资料，形成可检索知识库。",
            "农业知识图谱：将作物、政策、主体、市场、渠道、证据等组织为可视化知识树和图谱。",
            "智能问答：调用 Kimi/Moonshot 兼容接口或本地 Ollama 模型，生成面向农业经营的综合分析。",
            "实时搜索增强：根据用户问题接入网页搜索，补充最新政策、新闻、市场、气象和产业信息。",
            "证据链审计：支持默认隐藏、按需展开的来源引用、证据节点和知识树展示。",
            "报告导出：可根据当前问题生成 Word 报告，适用于项目申报、政策汇报和企业沟通。",
            "移动端适配：提供 PWA manifest 与 service worker，使手机浏览器可安装为轻应用。",
        ],
    )

    doc.add_heading("3. 源代码结构摘要", level=1)
    add_table(
        doc,
        ["目录", "说明"],
        [
            ["src/backend", "FastAPI 服务、问答接口、会话、模型提供商、联网搜索和区域情报逻辑。"],
            ["src/generation", "RAG 提示词、模型生成、回答格式控制和双赛道竞赛叙事。"],
            ["src/retrieval", "BM25、dense、hybrid、reranker 和检索路由模块。"],
            ["knowledge/ingestion", "文档抽取、分段、入库和知识库构建流程。"],
            ["web", "农业主题 GUI、问答面板、知识树、证据链、朗读和 PWA 入口。"],
            ["data/raw", "公开农业数据、本地政策、台湾农业数据样板和用户导入资料。"],
            ["docs/deliverables", "架构文档、使用说明、产品说明、申报材料和交付材料。"],
        ],
    )

    doc.add_heading("4. 软件特点", level=1)
    add_bullets(
        doc,
        [
            "面向农业场景的提示词和前端，而不是通用聊天界面。",
            "以证据对象和知识图谱组织资料，便于审计和项目申报。",
            "支持本地部署和联网增强，兼顾数据安全与实时性。",
            "可针对不同地区替换政策和数据包，具备复制推广能力。",
        ],
    )

    doc.add_heading("5. 登记材料准备清单", level=1)
    add_bullets(
        doc,
        [
            "软件著作权登记申请表。",
            "前后各连续 30 页源代码打印件或电子材料；不足 60 页则提交全部源代码。",
            "软件说明书或用户手册，本材料包中的使用说明书和产品说明书可作为基础。",
            "权利归属证明、开发完成日期、首次发表日期、申请人身份证明或企业营业执照。",
            "如有委托开发、合作开发或职务开发关系，需要准备对应协议或声明。",
        ],
    )

    path = OUT_DIR / "04_AgriKB_软件著作权登记材料_初稿.docx"
    save_doc(doc, path)
    write_markdown(
        OUT_DIR / "04_AgriKB_软件著作权登记材料_初稿.md",
        title,
        [
            ("软件信息", ["AgriKB 数智新农人经营中枢软件 V1.0。"]),
            ("模块", ["知识库构建、知识图谱、智能问答、实时搜索、证据链审计、报告导出、移动端适配。"]),
            ("材料", ["申请表、源代码、说明书、身份证明/营业执照、权属证明。"]),
        ],
    )
    return path


def build_deployment_guide() -> Path:
    title = "AgriKB 移动端与硬件部署适配说明"
    doc = Document()
    configure_document(doc, title)
    add_cover(doc, title, "手机、农业局工作站、合作社边缘盒子部署方案", "V1.0 交付稿")

    doc.add_heading("1. 部署形态选择", level=1)
    add_table(
        doc,
        ["形态", "推荐对象", "推荐硬件", "说明"],
        [
            ["手机/PWA", "农户、新农人、培训学员", "Android/iOS 手机 + 浏览器", "通过局域网访问服务端，添加到主屏幕使用。"],
            ["单机演示版", "比赛路演、农业局汇报", "Windows 笔记本，16GB 内存以上", "双击启动，使用本地知识库和模型回退。"],
            ["农业局工作站", "县域农业数据治理", "Windows 工作站，32GB 内存，1TB SSD", "适合长期运行、导入政策和台账。"],
            ["合作社边缘盒子", "田间/仓储/收购点", "迷你主机或工控机，16-32GB 内存", "可接入局域网摄像头、传感器和电子秤数据。"],
            ["云/私有化服务器", "多乡镇、多企业使用", "8 核 CPU、32GB+ 内存，可选 GPU", "支持多用户、统一数据更新和运维。"],
        ],
    )

    doc.add_heading("2. 手机端使用方式", level=1)
    add_numbered(
        doc,
        [
            "在电脑端启动 AgriKB 服务。",
            "查询电脑局域网 IP，例如 192.168.1.20。",
            "手机连接同一 Wi-Fi，访问 http://192.168.1.20:8010/ui/。",
            "浏览器菜单选择“添加到主屏幕”或“安装应用”，即可作为 PWA 使用。",
            "田间使用时建议打开联网搜索，结合当前天气、政策和市场信息生成建议。",
        ],
    )

    doc.add_heading("3. 数据接入建议", level=1)
    add_bullets(
        doc,
        [
            "气象：接入本地天气 API、气象站 CSV 或网页搜索结果，输出温度、降雨、风力、灾害预警和采收建议。",
            "地理环境：接入地块坐标、土壤、坡度、灌溉、遥感指数和 Google Earth Engine 农业数据目录。",
            "市场收购：接入农业局收购点、合作社报价、电商平台接口、批发市场公告和企业采购清单。",
            "政策新闻：接入政府网站、农业农村部门、句容本地政策文件和创业大赛通知。",
            "电商助农：预留抖音/快手/淘宝/京东/拼多多/微信小店/本地生活平台的商品、订单、库存和客服接口。",
        ],
    )

    doc.add_heading("4. 安全与运维", level=1)
    add_bullets(
        doc,
        [
            "API 密钥只放在 .env，不写入前端代码、说明书或公开材料。",
            "农业局正式部署建议使用内网、HTTPS、访问账号和日志审计。",
            "用户上传的台账和企业数据应按主体分库或分权限管理。",
            "定期备份 knowledge/chunks.db、knowledge/sessions、data/raw 和 .env。",
            "联网搜索结果属于实时情报，应在回答中标注时间，作为辅助判断而非唯一依据。",
        ],
    )

    doc.add_heading("5. 未来硬件扩展", level=1)
    add_bullets(
        doc,
        [
            "接入田间气象站：自动触发病虫害、采收和灌溉提醒。",
            "接入电子秤/收购机：形成收购台账、等级、价格和企业结算建议。",
            "接入摄像头/图像识别：识别病斑、成熟度、分级和包装质检。",
            "接入 NFC/二维码：形成从地块到电商订单的质量溯源链。",
        ],
    )

    path = OUT_DIR / "05_AgriKB_移动端与硬件部署适配说明.docx"
    save_doc(doc, path)
    write_markdown(
        OUT_DIR / "05_AgriKB_移动端与硬件部署适配说明.md",
        title,
        [
            ("手机端", ["同网访问服务端 IP:8010/ui/，可安装为 PWA。"]),
            ("硬件", ["比赛演示用笔记本；农业局用工作站；合作社用边缘盒子；规模化用私有化服务器。"]),
            ("扩展", ["气象站、电子秤、摄像头、二维码/NFC、电商平台接口。"]),
        ],
    )
    return path


def build_index(paths: list[Path]) -> None:
    index_lines = [
        "# AgriKB 数智新农人经营中枢交付材料包",
        "",
        f"- 生成日期：{TODAY.isoformat()}",
        "- 定位：2026 句容“福地青年英才”创业大赛，数智信息技术 × 新农人融合项目",
        "- 说明：专利和软著为初稿，正式提交前需要代理师/法务复核。",
        "",
        "## 文件清单",
        "",
    ]
    for path in paths:
        index_lines.append(f"- {path.name}")
    index_lines.extend(
        [
            "",
            "## 推荐使用顺序",
            "",
            "1. 比赛答辩：先看产品说明书和 PPT。",
            "2. 现场演示：按使用说明书启动并演示真实模型问答。",
            "3. 知识产权：专利交底书给代理师，软著材料用于准备登记。",
            "4. 落地部署：按移动端与硬件部署适配说明选择手机、工作站或边缘盒子。",
            "",
            "## 保密提醒",
            "",
            "API key、账号、企业台账和未公开政策材料不要写入公开 PPT 或公开申报附件。",
        ]
    )
    (OUT_DIR / "README_材料包索引.md").write_text("\n".join(index_lines), encoding="utf-8")


def verify_docx(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        document_xml = zf.read("word/document.xml")
    return {
        "file": path.name,
        "bytes": path.stat().st_size,
        "has_document_xml": "word/document.xml" in names,
        "xml_bytes": len(document_xml),
    }


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = [
        build_user_manual(),
        build_product_description(),
        build_patent_draft(),
        build_software_copyright(),
        build_deployment_guide(),
    ]
    build_index(paths)
    manifest = {
        "product": "AgriKB 数智新农人经营中枢",
        "date": TODAY.isoformat(),
        "output_dir": str(OUT_DIR),
        "docx": [verify_docx(path) for path in paths],
        "markdown": sorted(path.name for path in OUT_DIR.glob("*.md")),
    }
    (OUT_DIR / "materials_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
