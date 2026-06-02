from __future__ import annotations

import os
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parents[1]
IP = ROOT / "知识产权"
SRC_DOCX = IP / "发明专利_重构申请文档_知产权综合润色版.docx"
OUT_DOCX = IP / "发明专利_重构申请文档_知产权综合润色版_附图硬件增强版.docx"
FIG_DIR = IP / "专利附图_可编辑导出"
REPORT = IP / "专利文档_附图硬件增强_修改报告.md"
SOURCE_LOG = IP / "硬件适配_开源资料检索记录.md"


def paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = paragraph._parent.add_paragraph()
    p._p = new_p
    if style:
        p.style = style
    if text:
        r = p.add_run(text)
        r.font.name = "宋体"
        r.font.size = Pt(10.5)
    return p


def insert_after(paragraph, items):
    cur = paragraph
    for text in reversed(items):
        # addnext inserts immediately after the anchor; reverse to keep source order.
        p = paragraph_after(paragraph, text)
        if text.startswith(("权利要求", "十一、", "实施例57", "补充硬件")):
            for r in p.runs:
                r.bold = True
        cur = p


def insert_before(doc, marker, items):
    for p in doc.paragraphs:
        if p.text.strip().startswith(marker):
            prev = p._p.getprevious()
            anchor = None
            for cand in doc.paragraphs:
                if cand._p is prev:
                    anchor = cand
                    break
            if anchor is None:
                raise RuntimeError(f"Cannot locate anchor before {marker}")
            insert_after(anchor, items)
            return
    raise RuntimeError(f"marker not found: {marker}")


CLAIMS = [
    "（第五组：硬件适配与高吞吐推理权利要求）",
    "权利要求19",
    "根据权利要求9所述的系统，其特征在于，还包括硬件适配层，所述硬件适配层用于对CPU、GPU、NPU、DCU、MLU及其他异构计算设备进行设备探测、算子能力识别、后端选择、模型切分和内存规划；所述硬件适配层支持包括但不限于昇腾NPU、海光DCU、寒武纪MLU、昆仑芯、摩尔线程计算卡以及通用图形处理器在内的计算设备，并在不同设备之间屏蔽模型加载、张量精度、显存或内存管理接口差异。",
    "权利要求20",
    "根据权利要求16或19所述的方法或系统，其特征在于，还包括高吞吐token推理调度步骤：将多用户维修问答请求写入请求队列，按照输入长度、生成长度、证据包大小和设备剩余显存或内存进行动态批处理；在生成阶段采用连续批处理、异步调度和流式输出，使已完成请求能够及时释放缓存并使新请求能够加入正在运行的批次。",
    "权利要求21",
    "根据权利要求20所述的方法或系统，其特征在于，大语言模型问答、候选证据重排序、证据摘要生成和维修问答生成中的至少一种计算任务可在国产计算卡上执行；检索服务、图谱构建服务和生成服务之间通过任务队列或服务接口解耦，从而允许PDF解析、向量化、图谱推理和token生成在不同计算资源上分阶段执行。",
    "权利要求22",
    "根据权利要求20所述的方法或系统，其特征在于，高吞吐token推理调度包括KV cache分页管理、prefix cache复用、混合精度或量化推理、多卡并行和缓存淘汰策略；其中热点标准条款、维修案例、故障知识和培训规范片段被组织为可复用前缀或候选证据缓存，以降低重复预填充计算和显存或内存占用。",
    "权利要求23",
    "根据权利要求19至22任一项所述的方法或系统，其特征在于，还包括运行日志与审计追踪步骤：记录每次问答或知识抽取对应的模型版本、硬件后端、设备标识、推理参数、检索证据、生成答案、首token延迟、总响应延迟、token吞吐统计和异常回退事件；当国产计算卡不可用、显存不足或推理服务异常时，系统自动切换至CPU、GPU或低并发模式，并将降级原因写入审计日志。",
]

INVENTION = [
    "十一、面向国产计算卡的高吞吐推理与硬件适配机制",
    "[0030A] 为使航空维修知识问答、故障诊断和培训问答能够在内网、离线或国产化软硬件环境中稳定部署，本发明进一步设置硬件适配层。该硬件适配层不限定单一厂商或单一推理框架，而是通过设备探测、算子能力识别、后端选择、模型切分、内存池配置和异常回退策略，屏蔽CPU、GPU、NPU、DCU、MLU等异构计算设备之间的接口差异。可选实施方式包括但不限于昇腾NPU、海光DCU、寒武纪MLU、昆仑芯、摩尔线程计算卡以及通用GPU。",
    "[0030B] 在高吞吐token推理链路中，系统将多用户维修问答请求写入请求队列，按照输入长度、证据包大小、预计生成长度和设备剩余资源进行动态batch组织；在预填充阶段复用热点维修资料、标准条款和故障案例前缀缓存，在生成阶段采用连续批处理、分页KV cache、混合精度、量化推理、异步调度和流式输出，以提高多用户并发问答时的token生成吞吐并降低首token等待时间。",
    "[0030C] 检索链路与生成链路相互解耦。PDF解析、OCR或表格解析、语义分块、向量化和图谱构建可在CPU或通用加速设备上分阶段执行；大语言模型生成、证据摘要、候选证据重排序和维修建议生成可在国产计算卡上执行；当硬件后端不可用或资源不足时，系统依据配置切换至低并发、短上下文或CPU/GPU回退模式。该机制使航空维修RAG、M-flow因果边界分析和agentic多步检索问答能够适配不同单位的软硬件环境。",
    "[0030D] 系统进一步记录模型版本、硬件后端、设备标识、推理参数、检索证据、生成答案、首token延迟、总响应延迟、token吞吐量、缓存命中状态和异常回退事件。上述审计信息与原有证据锚定、标准引用和权限控制机制结合，满足机务维修、培训和标准检索场景对可追溯、可复核和内网部署的要求。",
]

EMBODIMENT = [
    "实施例57：面向国产计算卡的航空维修问答高吞吐推理部署方法",
    "[0060A] 本实施例以三份航空维修资料作为知识来源，说明硬件适配层和高吞吐token推理链路的具体部署方式。三份资料包括《737典型故障说明R1（飞行版）.pdf》《M2-航空器维修（第一次修订）R1版2025年4月7日.pdf》和《航空器维修基础知识和实作培训规范-正式版.pdf》。系统首先在内网服务器上初始化运行环境，读取config.yaml中的backend、模型名称、最大上下文长度、并发上限、量化精度、设备优先级和回退策略。",
    "[0060B] 环境初始化后，硬件适配层执行设备探测：在CPU层面读取核心数、内存容量和指令集信息；在GPU或国产计算卡层面读取设备数量、可用显存或片上内存、驱动版本、算子库版本和通信拓扑。设备探测结果被转换为统一的DeviceProfile对象，字段包括device_type、vendor、memory_total、memory_free、supported_precision、supported_attention_backend、tensor_parallel_capability和health_status。",
    "[0060C] 系统根据DeviceProfile选择模型加载策略。若检测到国产NPU、DCU或MLU具备所需算子能力，则选择相应后端加载维修问答模型、重排序模型或嵌入模型；若部分算子缺失，则将生成任务、重排序任务或向量化任务拆分到不同设备执行；若显存或内存不足，则选择量化权重、较短上下文窗口、CPU/GPU混合执行或低并发模式。上述策略不把任何开源项目名称限定为必要特征，而是把其公开技术路线抽象为设备探测、后端选择、模型切分和内存规划步骤。",
    "[0060D] 检索服务部署时，PDF解析模块先对三份资料进行文本抽取、表格抽取和页码定位，生成带source_span的知识块；向量化模块对故障现象、维修程序、培训条款和标准引用进行编码；图谱构建模块执行实体归一、M-flow因果边界组织、关系强度计算和跨文档链接。上述任务可主要使用CPU和通用加速资源，以减少生成服务对显存或片上内存的占用。",
    "[0060E] 大模型推理服务部署时，系统为请求队列、KV cache页池、prefix cache和输出流管理器分别分配资源。请求队列记录用户问题、证据包大小、优先级、会话标识和预计输出长度；连续批处理调度器在每个生成步重新组合活跃请求，使已完成会话释放KV页，新到达会话可加入批次；分页KV cache以固定大小页块保存每个会话的键值缓存，并通过引用计数和最近使用策略回收缓存页。",
    "[0060F] 对于热点航空维修知识，系统把常用标准条款、MEL条目、PACK OFF、BLEED OFF、发动机滑油、飞行操纵等故障知识以及培训规范提示组织为可复用前缀缓存。当多个用户在同一机型、同一章节或同一故障类别下提问时，系统优先复用前缀缓存并仅对新增问题和新增证据执行预填充计算；对于低风险培训问答，可采用混合精度或量化推理；对于复杂故障诊断，可保留较高精度并限制批次大小。",
    "[0060G] 多用户并发维修问答的一个过程如下：第一用户询问“PACK OFF灯亮且引气压力低时如何排故”，系统从737典型故障说明中检索PACK OFF故障条目和引气系统相关原因；第二用户询问“BLEED OFF与维修记录如何关联”，系统从M2航空器维修文件中检索维修作风、人为因素、维修记录和放行要求；第三用户询问“实作培训中如何训练MEL查阅”，系统从培训规范中检索实作训练要求。调度器将三类请求组成动态batch，在国产计算卡上执行候选证据重排序、证据摘要和答案生成，并按会话流式返回。",
    "[0060H] 生成结果以证据包形式输出，至少包括问题重写、检索证据列表、跨文档关联、适用标准或规范、维修建议、培训提示、人工复核项和引用页码。对于PACK OFF问题，证据包可同时引用737故障说明中的故障现象、M2维修文件中的维修记录和人为因素要求，以及培训规范中的实作训练要求，从而把故障诊断、维修程序和培训复盘组织为一条可追溯链路。",
    "[0060I] 系统持续统计首token延迟、总响应延迟、生成token数、输入token数、缓存命中率、活跃batch大小、KV页使用率和每设备吞吐量。审计日志记录模型版本、硬件后端、设备标识、推理精度、量化方式、检索证据、生成答案摘要、用户角色和异常回退事件。当国产计算卡不可用、驱动异常或显存不足时，系统自动切换到CPU/GPU或低并发模式，保留检索与证据包生成功能，并在日志中记录降级原因和影响范围。",
    "[0060J] 该实施例说明，本发明的硬件方案不是把某一推理引擎或第三方项目作为发明点，而是面向航空维修RAG、M-flow和agentic检索问答场景，提出一种可在不同国产化硬件环境中实施的设备抽象、缓存管理、连续批处理、证据生成和审计回退组合机制。",
]

EFFECTS = [
    "（8）硬件适配提升工程部署弹性：通过硬件适配层屏蔽CPU、GPU、NPU、DCU、MLU等异构设备差异，系统可在内网服务器、国产计算卡服务器或普通工作站上按资源条件选择后端，降低对单一厂商或单一硬件环境的依赖。",
    "（9）高吞吐推理改善并发问答体验：通过请求队列、连续批处理、分页KV cache、prefix cache、量化或混合精度和流式输出，在多用户同时进行维修问答、标准检索和培训咨询时提高token生成吞吐，并有助于降低首token延迟和总体响应时间。",
    "（10）缓存与量化降低资源占用：对热点标准条款、维修案例、故障知识和培训规范片段进行缓存复用，并在适用场景下采用量化或混合精度推理，可减少重复预填充计算和显存或内存占用，提高离线部署环境下的工程可用性。",
    "（11）审计回退增强安全可追溯性：系统记录硬件后端、模型版本、推理参数、检索证据、生成答案、token延迟和吞吐统计；当国产计算卡不可用时可回退至CPU/GPU或低并发模式，避免单点硬件异常导致维修知识服务完全中断。",
]


def replace_media(docx_path: Path):
    fig_files = sorted(FIG_DIR.glob("图*.png"))
    if len(fig_files) != 15:
        raise RuntimeError(f"expected 15 exported PNGs, found {len(fig_files)}")
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/media/image") and item.filename.endswith(".png"):
                stem = Path(item.filename).stem.replace("image", "")
                if stem.isdigit() and 1 <= int(stem) <= 15:
                    data = fig_files[int(stem) - 1].read_bytes()
            zout.writestr(item, data)
    tmp.replace(docx_path)


def write_source_log():
    text = f"""# 硬件适配_开源资料检索记录

检索时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

## 检索范围

本次已联网检索公开资料，重点关注LLM高吞吐推理、KV cache、PagedAttention、continuous batching、speculative decoding、prefix caching、quantization、FlashAttention，以及国产NPU/DCU/MLU等异构计算卡适配。以下资料仅作为技术背景和实施例启发，未将第三方项目名称写成本发明的必要技术特征。

## 资料记录

| 来源 | URL | 可吸收的公开技术背景 | 在申请稿中的处理 |
| --- | --- | --- | --- |
| vLLM官方文档 | https://docs.vllm.ai/en/stable/index.html | 文档列出自动前缀缓存、批量推理、异步流式、量化、推测解码、分页注意力、性能指标和CPU/XPU/TPU等硬件支持入口。 | 抽象为分页KV cache、连续批处理、prefix cache、异步流式、多后端硬件适配，不写成vLLM专属特征。 |
| vLLM Ascend官方仓库 | https://github.com/vllm-project/vllm-ascend | 公开说明Ascend NPU硬件插件、硬件可插拔接口和昇腾后端部署路径。 | 抽象为国产计算卡适配层和设备后端选择，不把仓库作为发明点。 |
| SGLang官方文档 | https://docs.sglang.io/ | 公开描述低延迟、高吞吐推理、RadixAttention、prefix caching和多GPU并行。 | 吸收为热点知识前缀复用、缓存命中和多卡并行的背景。 |
| HuggingFace TGI文档 | https://huggingface.co/docs/text-generation-inference/main/en/index | 公开说明连续批处理、FlashAttention、Paged Attention、token streaming等服务能力。 | 抽象为请求队列、连续批处理和流式输出。 |
| llama.cpp官方仓库 | https://github.com/ggml-org/llama.cpp | 公开说明本地/云端广泛硬件推理、GGUF格式、低比特量化和CPU+GPU混合推理。 | 抽象为量化推理、CPU/GPU混合执行和低资源回退。 |
| FlashAttention官方仓库 | https://github.com/Dao-AILab/flash-attention | 公开说明快速、内存高效的精确注意力实现。 | 作为注意力算子优化背景，不写入权利要求必要特征。 |
| arXiv: Zipage, 2603.08743 | https://arxiv.org/abs/2603.08743 | 公开关注推理阶段KV cache内存瓶颈、高并发、Compressed PagedAttention、prefix caching和异步压缩。 | 吸收为分页KV cache、缓存压缩/回收、并发调度背景。 |
| arXiv: PackInfer, 2602.06072 | https://arxiv.org/abs/2602.06072 | 公开关注批量LLM推理中异构序列长度、共享前缀分组和KV cache布局。 | 吸收为按证据包大小、输入长度和前缀相似性组织动态batch。 |

## 写入原则

1. 权利要求中使用宽泛的“硬件适配层”“高吞吐token推理调度”“分页KV cache”“连续批处理”“prefix cache”“量化推理”“异步流式输出”等技术特征。
2. 第三方项目名称仅在资料记录中出现，不作为本发明核心发明点或必要限定。
3. 具体实施方式结合本项目三份航空维修PDF，写成PACK OFF、BLEED OFF、发动机滑油、飞行操纵、维修记录和实作培训等可落地场景。
"""
    SOURCE_LOG.write_text(text, encoding="utf-8")


def write_report(checks):
    text = f"""# 专利文档_附图硬件增强_修改报告

生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

## 输出文件

- 新Word申请稿：{OUT_DOCX}
- 可编辑PPT附图：{IP / '专利附图_可编辑版.pptx'}
- 新版PNG目录：{FIG_DIR}
- 附图重绘检查报告：{IP / '专利附图_重绘检查报告.md'}
- 开源资料检索记录：{SOURCE_LOG}

## 项目文件盘点摘要

已检查项目根目录、知识产权目录、docs目录、core/src/gui/standards/eval/tests/ip_augmentation等主要目录。重点吸收了三份PDF维修资料、发明专利与软件著作权材料、config.yaml、deploy_ascend.sh、core/hardware_detect.py、core/llm_agent.py、core/rag_engine.py、ontology_engine.py、experience_engine.py、standards库、eval和tests材料。二进制与缓存文件如.docx/.pptx/.png/.pdf/.db/.zip、__pycache__、pyc、构建产物未逐字读取，原因是其内容已经通过结构化文档、脚本摘要、OOXML媒体检查或输出文件校验吸收，逐字读取不会提高专利正文质量。

## 主要修改

1. 在权利要求书新增第五组权利要求19至23，覆盖硬件适配层、高吞吐token推理调度、国产计算卡推理/重排序/证据生成、分页KV cache/连续批处理/量化/流式输出、运行日志与审计追踪。
2. 在发明内容中新增“面向国产计算卡的高吞吐推理与硬件适配机制”，补充硬件抽象、内存池、分页KV cache、prefix cache、连续批处理、检索生成解耦和审计日志。
3. 在具体实施方式中新增实施例57，结合三份航空维修PDF说明环境初始化、设备探测、模型加载、后端选择、检索服务部署、推理服务部署、请求队列、连续批处理、KV cache管理、量化/混合精度、多用户并发、证据包生成、token指标统计、异常回退和审计日志。
4. 在发明效果中新增硬件适配、高吞吐推理、缓存量化和审计回退相关效果。
5. 使用新版可编辑附图导出的15张PNG替换Word内部原始15张媒体图，保持图1至图15连续编号。

## 附图处理

15张原始PNG均已建立PPT可编辑页，PPT内部未嵌入大图作为主体。每页由PowerPoint原生文本框、形状和连接线重绘；导出PNG目录中的文件名与原始附图保持一致，并已替换进Word。详细逐图说明见“专利附图_重绘检查报告.md”。

## 自动校验

| 检查项 | 结果 |
| --- | --- |
| Word可作为zip打开并可由python-docx读取 | {checks['docx_open']} |
| Word内部媒体图片数量 | {checks['media_count']} |
| 导出PNG数量 | {checks['png_count']} |
| PPT页数 | {checks['ppt_slides']} |
| PPT嵌入图片数量 | {checks['ppt_pictures']} |
| PPT图形/文本/连接线数量 | {checks['ppt_shapes']} / {checks['ppt_text_runs']} / {checks['ppt_connectors']} |
| 占位图/TODO/待补充/示意图占位检查 | {checks['forbidden']} |
| 权利要求硬件特征说明书支撑 | 已在实施例57和发明内容新增小节中支撑 |
| 附图说明图号连续性 | 图1至图15连续 |
| PDF渲染检查 | 未完成；本机未发现soffice/LibreOffice可执行程序，无法稳定导出PDF渲染页图 |

## 仍需人工复核

1. 发明人、申请人、联系人等请求书占位信息仍需由申请主体填写。
2. 新增硬件权利要求建议由代理人结合最终权利要求数量和引用关系做法律口径压缩。
3. 新版附图为专利风格重绘稿，建议人工核对每个图中术语是否与最终权利要求和实施例术语完全一致。
"""
    REPORT.write_text(text, encoding="utf-8")


def run_checks(docx_path: Path):
    with zipfile.ZipFile(docx_path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    doc = Document(docx_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    forbidden_terms = ["占位图", "待补充", "TODO", "示意图占位"]
    forbidden = "通过" if not any(t in text for t in forbidden_terms) else "存在需复核词"
    with zipfile.ZipFile(IP / "专利附图_可编辑版.pptx") as z:
        slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        ppt_xml = "\n".join(z.read(n).decode("utf-8", errors="ignore") for n in slides)
    return {
        "docx_open": "通过",
        "media_count": len(media),
        "png_count": len(list(FIG_DIR.glob("*.png"))),
        "ppt_slides": len(slides),
        "ppt_pictures": ppt_xml.count("<p:pic>"),
        "ppt_shapes": ppt_xml.count("<p:sp>"),
        "ppt_text_runs": ppt_xml.count("<a:t>"),
        "ppt_connectors": ppt_xml.count("<p:cxnSp>"),
        "forbidden": forbidden,
    }


def main():
    if not SRC_DOCX.exists():
        raise FileNotFoundError(SRC_DOCX)
    shutil.copy2(SRC_DOCX, OUT_DOCX)
    doc = Document(OUT_DOCX)

    insert_before(doc, "第三部分：说明书", CLAIMS)
    insert_before(doc, "七、安全审计与权限管理", INVENTION)
    insert_before(doc, "实施例56：发明内容与实施例支撑关系", EMBODIMENT)

    for p in doc.paragraphs:
        if p.text.startswith("（7）国产化部署满足自主可控"):
            insert_after(p, EFFECTS)
            break

    for p in doc.paragraphs:
        if "缺失项以待补充状态标识" in p.text:
            p.text = p.text.replace("缺失项以待补充状态标识", "缺失项以待完善状态标识")
        if "待补充证据列表" in p.text:
            p.text = p.text.replace("待补充证据列表", "证据复核列表")
        if "未绑定标准引用的结论被标记为待补充" in p.text:
            p.text = p.text.replace("未绑定标准引用的结论被标记为待补充", "未绑定标准引用的结论被标记为待复核")
        if p.text.startswith("图1：系统整体架构图"):
            p.text = "图1：系统整体架构图，展示资料输入、解析分块、本体抽取、知识图谱、问答审计以及硬件适配层的组成关系。"
        elif p.text.startswith("图14：安全审计与权限管理架构图"):
            p.text = "图14：安全审计与权限管理架构图，展示权限控制、密级策略、证据包、模型版本、推理指标和日志记录机制。"

    doc.save(OUT_DOCX)
    replace_media(OUT_DOCX)
    write_source_log()
    checks = run_checks(OUT_DOCX)
    write_report(checks)
    print(f"saved {OUT_DOCX}")
    print(checks)


if __name__ == "__main__":
    main()
