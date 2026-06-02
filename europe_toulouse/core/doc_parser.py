"""
多格式文档解析器
支持: PDF, DOCX, XLSX, PPTX, TXT, MD, HTML, XML, JSON, YAML, CSV, RST, TEX
"""
import os
import json
import csv
import io
import re
from pathlib import Path
from typing import Optional

try:
    import chardet
except ImportError:  # pragma: no cover - optional dependency fallback
    chardet = None


def detect_encoding(file_path: str) -> str:
    with open(file_path, "rb") as f:
        raw = f.read(10000)
    if chardet is None:
        for candidate in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                raw.decode(candidate)
                return candidate
            except UnicodeDecodeError:
                continue
        return "utf-8"
    result = chardet.detect(raw)
    return result.get("encoding", "utf-8") or "utf-8"


def parse_pdf(file_path: str) -> str:
    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(file_path)
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if text.strip():
            return text
    except Exception:
        pass

    import pdfplumber

    texts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                texts.append(t)
    return "\n\n".join(texts)


def parse_docx(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(c.text for c in row.cells))
    return "\n".join(paragraphs)


def parse_pptx(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    texts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_text = [f"--- Slide {i} ---"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text)
        texts.append("\n".join(slide_text))
    return "\n\n".join(texts)


def parse_xlsx(file_path: str) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True, data_only=True)
    texts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        texts.append(f"=== Sheet: {sheet_name} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                texts.append(" | ".join(cells))
    return "\n".join(texts)


def parse_csv_file(file_path: str) -> str:
    enc = detect_encoding(file_path)
    with open(file_path, "r", encoding=enc, errors="replace") as f:
        reader = csv.reader(f)
        rows = [" | ".join(row) for row in reader]
    return "\n".join(rows)


def parse_html(file_path: str) -> str:
    from bs4 import BeautifulSoup
    enc = detect_encoding(file_path)
    with open(file_path, "r", encoding=enc, errors="replace") as f:
        soup = BeautifulSoup(f.read(), "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


def parse_xml(file_path: str) -> str:
    from lxml import etree
    tree = etree.parse(file_path)
    return etree.tostring(tree, pretty_print=True, encoding="unicode")


def parse_json_file(file_path: str) -> str:
    enc = detect_encoding(file_path)
    with open(file_path, "r", encoding=enc, errors="replace") as f:
        data = json.load(f)
    return json.dumps(data, indent=2, ensure_ascii=False)


def parse_yaml_file(file_path: str) -> str:
    import yaml
    enc = detect_encoding(file_path)
    with open(file_path, "r", encoding=enc, errors="replace") as f:
        data = yaml.safe_load(f)
    return json.dumps(data, indent=2, ensure_ascii=False)


def parse_text(file_path: str) -> str:
    enc = detect_encoding(file_path)
    with open(file_path, "r", encoding=enc, errors="replace") as f:
        return f.read()


PARSER_MAP = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".doc": parse_docx,
    ".pptx": parse_pptx,
    ".xlsx": parse_xlsx,
    ".xls": parse_xlsx,
    ".csv": parse_csv_file,
    ".html": parse_html,
    ".htm": parse_html,
    ".xml": parse_xml,
    ".json": parse_json_file,
    ".yaml": parse_yaml_file,
    ".yml": parse_yaml_file,
    ".txt": parse_text,
    ".md": parse_text,
    ".rst": parse_text,
    ".tex": parse_text,
}


def parse_document(file_path: str) -> Optional[dict]:
    """解析任意支持格式的文档，返回 {path, name, ext, content, size_kb}"""
    p = Path(file_path)
    ext = p.suffix.lower()
    parser = PARSER_MAP.get(ext)
    if parser is None:
        return None
    try:
        content = parser(str(p))
        return {
            "path": str(p.resolve()),
            "name": p.name,
            "ext": ext,
            "content": content,
            "size_kb": round(p.stat().st_size / 1024, 1),
        }
    except Exception as e:
        return {
            "path": str(p.resolve()),
            "name": p.name,
            "ext": ext,
            "content": f"[解析失败: {e}]",
            "size_kb": round(p.stat().st_size / 1024, 1),
        }


def scan_directory(directory: str, supported_exts: list = None) -> list:
    """递归扫描目录，返回所有可解析文档信息列表"""
    if supported_exts is None:
        supported_exts = list(PARSER_MAP.keys())
    results = []
    for root, _, files in os.walk(directory):
        for fname in files:
            fpath = os.path.join(root, fname)
            ext = Path(fpath).suffix.lower()
            if ext in supported_exts:
                doc = parse_document(fpath)
                if doc:
                    results.append(doc)
    return results
