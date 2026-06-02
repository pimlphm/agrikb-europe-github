from __future__ import annotations

import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement


IP = Path(__file__).resolve().parent
DOCX = IP / "发明专利_重构申请文档_知产权综合润色版_附图硬件增强版.docx"
BACKUP = IP / f"发明专利_重构申请文档_知产权综合润色版_附图硬件增强版_四实施例重写前备份_{datetime.now():%Y%m%d_%H%M%S}.docx"
REPORT = IP / "专利文档_附图硬件增强_修改报告.md"


def find_para(doc: Document, prefix: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i
    raise RuntimeError(f"marker not found: {prefix}")


def remove_para(p):
    p._element.getparent().remove(p._element)


def add_after(anchor, text: str, bold: bool = False):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    p = anchor._parent.add_paragraph()
    p._p = new_p
    r = p.add_run(text)
    r.font.name = "宋体"
    r.font.size = Pt(10.5)
    r.bold = bold
    return p


IMPLEMENTATION = [
    ("导言", [
        "以下结合附图和实施例，对本发明作进一步详细说明。为避免将同一技术链路拆分为大量零散片段，以下实施例按照工程实现顺序合并为六个大的实施例。每个实施例均给出输入对象、处理步骤、数据结构、判定规则、输出结果和异常回退方式，所属领域技术人员可据此在不依赖特定软件名称或特定硬件厂商的情况下实施本发明。",
    ]),
    ("实施例1：系统初始化、部署环境和基础数据对象", [
        "[0030] 本实施例说明系统从空白部署环境到可接收航空维修资料的初始化过程。系统至少包括资料接入模块、文档解析模块、语义分块模块、本体知识图谱模块、标准条款检索模块、多角色模型推理模块、证据锚定生成模块、经验库模块、审计日志模块、图形复核模块和硬件适配模块。上述模块可部署在同一工作站、内网服务器或私有化服务集群中，也可按检索服务、图谱服务和生成服务拆分部署。",
        "[0031] 初始化步骤包括：S101，读取配置文件，获得资料目录、输出目录、模型后端、最大上下文长度、检索top_k、稀疏检索权重、稠密检索权重、重排序候选数量、审计日志路径、特征开关和硬件优先级；S102，创建raw、processed、chunk、graph、standard、experience、audit和export目录；S103，加载标准条款库、实体类型表、关系类型表、关联强度权重表、提示词模板和权限策略；S104，执行硬件探测，形成DeviceProfile对象；S105，启动检索索引、图谱存储和模型服务健康检查。",
        "[0032] DeviceProfile对象至少包括device_type、vendor、device_id、memory_total、memory_free、driver_version、backend_name、supported_precision、supported_attention_backend、tensor_parallel_capability、health_status和fallback_priority字段。若未检测到可用加速卡，device_type被设为CPU，系统自动启用低并发、短上下文和保守生成策略；若检测到多个设备，系统按设备健康状态、剩余内存和配置优先级选择生成后端、向量化后端和重排序后端。",
        "[0033] 系统为每个输入文件创建DocumentRecord对象，字段包括document_id、file_name、file_type、source_path、sha1或等效指纹、version_tag、security_level、import_time、parser_status、page_count、language、domain_tags和owner。后续所有知识块、实体、关系、答案和审计记录均引用document_id，不直接依赖本地文件路径作为业务标识，从而支持迁移部署和版本回滚。",
        "[0034] 系统的核心中间数据为ChunkRecord、EntityRecord、RelationRecord、EvidenceRecord和AuditRecord。ChunkRecord保存chunk_id、document_id、page_no、section_title、chunk_type、text、table_meta、source_span和embedding_status；EntityRecord保存entity_id、name、entity_type、ontology_layer、properties、source_span、standard_refs和confidence；RelationRecord保存relation_id、head_id、tail_id、relation_type、strength、color_level、evidence_ids、standard_refs和audit_status；EvidenceRecord保存evidence_id、document_id、page_no、quote_summary、source_span、standard_refs和reliability；AuditRecord保存operation_id、operator、time、model_version、hardware_backend、parameters、input_summary、output_summary、latency和fallback_event。",
        "[0035] 初始化完成的判定条件为：标准库可查询，资料目录可写入，图谱存储可创建节点和边，检索索引可接受新增chunk，模型后端返回健康状态或启用离线回退模式，审计日志可追加写入。若任一条件失败，系统不进入正式知识抽取流程，而是输出初始化异常报告，列出失败模块、失败原因、回退策略和人工处理建议。"
    ]),
    ("实施例2：航空维修资料解析、语义分块和本体约束知识抽取", [
        "[0036] 本实施例说明如何把PDF、Word、Markdown、文本和表格资料转换为可抽取的结构化知识。资料接入模块接收输入文件后，首先根据文件类型选择解析器；对于PDF，优先提取页码、正文、表格、标题层级、页眉页脚和图片说明；对于Word或Markdown，提取段落、标题、表格和列表；对于扫描质量较差的页面，可标记为需要OCR或人工复核。解析结果统一写入PageRecord和BlockRecord。",
        "[0037] 语义分块步骤包括：S201，按章节标题、页码、表格边界和段落长度生成初始块；S202，识别块类型，至少包括故障现象块、维修动作块、标准条款块、培训要求块、安全评估块、追溯证据块和一般说明块；S203，对超过长度阈值的块按句子边界继续拆分，对过短且同属一页一节的块进行合并；S204，为每个块生成source_span，记录document_id、page_no、section_title、paragraph_index、char_start和char_end；S205，写入ChunkRecord并进入候选抽取队列。",
        "[0038] 本体约束抽取按照三层本体执行。顶层类型包括实体、过程、功能、质量、角色和倾向；核心层类型包括需求、设计、验证、接口、约束、规范、组件、系统、测试、风险、标准、文档、计划、报告、决策、指标、工具和人员；航空维修领域层类型至少包括故障现象、失效状态、机载系统、部件、维修动作、MEL条目、放行限制、培训科目、人为因素、维修记录、安全目标、DAL分配和证据片段。抽取时，模型或规则引擎只能从允许类型中选择实体类型，不能自由生成未登记类型。",
        "[0039] 实体抽取流程为：S211，从ChunkRecord中识别候选术语；S212，使用术语词典、标准编号规则、英文缩写规则和上下文触发词判断候选类型；S213，将候选实体写成EntityRecord；S214，若同一document_id和相近source_span内出现同名实体，则合并为同一局部实体；S215，若跨文档出现相似实体，则暂存为candidate_same_as关系，等待实施例4中的跨文档共指消解确认。",
        "[0040] 关系抽取流程为：S221，根据关系触发词、句法方向、表格列名和章节语义生成候选关系；S222，根据关系类型允许的头尾实体类型过滤不合法关系，例如failure_leads_to关系的头实体应为失效状态或故障原因，tail实体应为故障现象、危险、功能丧失或维修后果；S223，计算关联强度，关联强度由关系类型权重、证据距离权重、标准条款权重、图结构权重和模型置信度加权得到；S224，依据强度区间映射为七级显示等级；S225，将通过阈值的关系写入RelationRecord，将未通过阈值但证据存在的关系写入候选池。",
        "[0041] 本实施例的输出为一个可追溯的初始知识图谱：每个节点均有来源页码和证据跨度，每条边均有关系类型、强度、标准引用和候选/正式状态。若解析器无法提取可靠文本，系统将相应页面标记为low_quality_page；若实体类型无法确定，标记为unknown_candidate而不写入正式图谱；若模型输出不符合预设JSON或字段缺失，进入格式回退流程，由规则抽取器生成保守候选结果并写入审计日志。"
    ]),
    ("实施例3：标准条款注入、多角色协同、六层校验和经验闭环", [
        "[0042] 本实施例说明知识抽取结果如何与适航标准、维修规范和工程经验结合。系统先把CAAC、FAA、EASA以及维修培训规范中的条款对象化，每个StandardClause对象至少包括standard_id、authority、version、clause_id、title_cn、title_en、scope、keywords、applicable_entity_types、required_fields、equivalent_refs、summary和effective_date。条款对象不要求保存全文，至少保存可检索摘要、适用实体类型和必填字段约束。",
        "[0043] 标准条款检索采用稀疏检索、稠密检索和类型匹配三路融合。S301，系统从当前ChunkRecord和EntityRecord中抽取查询词，包括故障码、系统名、部件名、MEL编号、标准编号和动作词；S302，BM25路径检索标准编号、中文标题、英文标题和关键词；S303，向量路径检索语义相近条款；S304，类型匹配路径判断条款适用实体类型是否覆盖当前实体类型；S305，按R(q,c)=λ1·BM25+λ2·CosSim+λ3·TypeMatch+λ4·ProjectContext计算得分；S306，取Top-K条款作为标准上下文注入抽取和生成流程。",
        "[0044] 多角色协同流程包括知识提取器、本体推理器、安全评估器、追溯检查器、变更影响分析器和文档整理器。知识提取器负责候选实体和关系；本体推理器负责端点类型、层级归属和潜在链接；安全评估器负责危险、失效状态、安全目标、DAL建议和缓解措施；追溯检查器负责需求、设计、测试、维修记录和标准条款之间的覆盖关系；变更影响分析器沿图谱边输出上游、下游、横向、安全和培训影响；文档整理器只对已经通过校验的对象生成证据包和报告。",
        "[0045] 六层校验按固定顺序执行。第一层为格式校验，检查JSON字段、类型、必填项和枚举值；第二层为本体一致性校验，检查实体类型和关系端点是否匹配；第三层为标准引用校验，凡涉及安全目标、DAL、MEL、放行限制、维修记录和培训要求的对象必须绑定标准、规范或证据片段；第四层为证据跨度校验，检查source_span是否能回指到具体文档页码和段落；第五层为冲突检测，比较不同角色或不同文档对同一实体属性、原因关系或处置建议的差异；第六层为置信度阈值校验，低于阈值的对象进入人工复核队列。",
        "[0046] 经验闭环在校验后执行。系统将通过校验的故障分析、维修处置、标准引用、人工复核意见和审计记录组织为ExperienceEntry，字段包括experience_id、title、aircraft_type、task_type、applicable_conditions、content、evidence_ids、standard_refs、confidence、review_status、created_by和usage_count。经验条目先进入临时经验库，经工程人员审核后可提升为永久经验。后续遇到相同机型、相同系统、相似故障或相同标准条款时，系统把匹配经验作为上下文注入，但经验只作为辅助约束，不替代标准条款和原文证据。",
        "[0047] 本实施例的输出包括通过校验的正式知识图谱、候选池、冲突清单、标准引用清单、经验条目和审计日志。若标准库缺少对应条款，系统记录standard_gap事件并提示人工补录；若不同角色结论冲突，系统不自动选择一方，而是保留双方证据、置信度和来源页码；若经验条目未经审核，系统仅在回答中作为低权重背景使用。"
    ]),
    ("实施例4：围绕三份航空维修PDF的端到端故障知识图谱构建", [
        "[0048] 本实施例以三份实际航空维修资料为输入，说明从PDF资料到故障知识图谱、故障树、因果树和跨文档证据包的完整实施过程。输入资料包括《737典型故障说明R1（飞行版）.pdf》《M2-航空器维修（第一次修订）R1版2025年4月7日.pdf》和《航空器维修基础知识和实作培训规范-正式版.pdf》。三份资料分别提供故障案例、维修管理和培训实作规范，系统将其作为同一航空维修知识域的不同证据源处理。",
        "[0049] 对《737典型故障说明R1（飞行版）.pdf》，系统执行以下步骤：S401，解析故障目录和故障详情页，识别PACK OFF、BLEED OFF、DUAL BLEED OFF、发动机滑油、飞行操纵和起飞构型等故障主题；S402，为每个故障主题建立Episode对象；S403，在Episode下建立故障现象、运行条件、可能原因、排故动作、排除原因、维修建议和预防提示等Facet；S404，把每个Facet中的具体表述转化为FacetPoint并关联EntityRecord；S405，依据触发词和章节语义建立symptom_of、caused_by、requires_check、excluded_cause和maintenance_action等关系。",
        "[0050] 以PACK OFF为例，系统把“PACK OFF灯亮”“引气压力异常”“空调组件工作异常”“查阅MEL”“完成维修记录”分别抽取为故障现象、失效状态、部件状态、标准条款和维修动作实体。若同一段落同时出现“压力低”和“组件关断”，系统不直接将二者互为因果，而是检查上下文是否存在“导致、引起、由于、检查、排除”等触发词，并结合MEL或维修程序引用确定关系方向。最终输出的PACK OFF子图至少包含故障现象节点、引气系统节点、空调组件节点、MEL条目节点、维修动作节点和培训提示节点。",
        "[0051] 对《M2-航空器维修（第一次修订）R1版2025年4月7日.pdf》，系统重点抽取维修作风、人为因素、维修记录、放行要求、工具管理、工卡执行和诚信管理等实体。S411，将规章或程序性表述归为maintenance_rule或procedure_requirement；S412，将人为因素、疲劳、沟通、交接、复核等内容归为human_factor；S413，将维修记录、签署、放行和复查归为record_requirement；S414，将这些对象与737故障案例中的维修动作、MEL条目和复核要求建立supports或requires_record关系。",
        "[0052] 对《航空器维修基础知识和实作培训规范-正式版.pdf》，系统重点抽取培训科目、实作步骤、考核要点、安全注意事项和MEL查阅训练要求。S421，将培训目标、训练项目、实作条件和考核标准对象化；S422，将其与故障案例中的PACK OFF、BLEED OFF、发动机滑油和飞行操纵故障建立training_supports关系；S423，当培训规范要求实作演练或风险提示时，系统在故障知识图谱中增加training_tip节点，供问答生成和报告输出引用。",
        "[0053] 跨文档关联分为共指消解、标准桥接和图结构确认三步。共指消解统一“空调组件”“PACK组件”“组件关断”等近义表达；标准桥接把MEL条目、维修记录要求和培训规范条款连接到相同或相关的标准对象；图结构确认检查两个实体是否共享部件、系统、动作或标准引用。只有同时满足名称相似、类型兼容和证据支持的关联才写入正式图谱，其他关联保留为候选关系。",
        "[0054] 故障树生成以顶事件为入口。系统接收顶事件名称，例如“引气系统压力低”或“PACK OFF灯亮”，在图谱中定位对应failure_state节点；沿caused_by、failure_leads_to和requires_check关系向上展开原因节点；将可直接检查或可维修的部件、传感器、活门、管路和控制逻辑标记为基本事件；根据多个原因是否独立、共同作用或排除关系推断OR门、AND门或NOT标记。生成结果包括顶事件、中间事件、基本事件、逻辑门类型、证据页码和人工复核标志。",
        "[0055] 因果树生成以时间和证据强度为主线。系统从故障现象开始，沿confirmed优先、possible次之、candidate再次之的顺序组织因果链；若同一原因链同时被737故障说明和M2维修文件支持，则提高关联强度；若培训规范只提供操作提示而不支持因果判断，则仅作为training_tip挂接，不进入因果主链。由此可避免把培训要求误当作故障原因，也可避免把维修动作误当作故障现象。",
        "[0056] 本实施例输出包括图7所示PACK OFF子图、图8所示故障树、图9所示因果树、图10所示跨文档关联网络，以及面向用户问答的证据包。证据包至少包括故障主题、实体清单、关系清单、标准或规范引用、来源页码、维修建议、培训提示、不可自动确认事项和审计编号。若某一PDF页解析质量不足或来源证据冲突，系统在证据包中标注需要人工复核，而不是生成确定性维修结论。"
    ]),
    ("实施例5：检索增强问答、图形界面复核、报告输出和审计回退", [
        "[0057] 本实施例说明用户提出航空维修问题后，系统如何检索、生成、复核和审计。用户问题可以是事实型问题、故障诊断问题、标准检索问题、培训问答问题或跨文档分析问题。系统首先创建QueryRecord，字段包括query_id、user_id、query_text、query_type、aircraft_type、priority、created_time、route、status和audit_id。",
        "[0058] 查询路由规则为：若问题包含明确故障码、标准编号、部件名或MEL条目，则优先使用基线混合检索路径；若问题包含“为什么、如何排故、影响哪些、关联哪些标准”等多步推理表达，则使用Agentic多步路径；若问题要求“综合三份资料、列出所有证据、比较多个故障链路”，则使用递归检索路径；若高级路径未启用或模型不可用，则自动降级到基线混合检索路径，并在AuditRecord中记录fallback_event。",
        "[0059] 混合检索路径包括：S501，对用户问题进行分词、标准编号识别、实体识别和同义词扩展；S502，BM25检索命中精确术语、故障码、MEL编号和页码附近片段；S503，稠密向量检索命中语义相近的维修程序、培训要求和案例描述；S504，将两路结果按归一化分数和RRF融合；S505，重排序器根据关键词重叠、实体类型、来源文档多样性、证据页码完整性和标准引用完整性重新排序；S506，取Top-N证据生成EvidenceBundle。",
        "[0060] Agentic多步路径在每一步保存step_id、current_query、retrieved_evidence、decision、next_query和stop_reason。第一步检索直接证据，第二步根据缺口补充标准或培训证据，第三步检查是否存在冲突或缺失来源，第四步生成最终证据包。若某一步证据不足，模型只能改写查询或扩大检索范围，不能凭空补充事实；若达到最大步数仍不足，系统输出证据不足说明和人工复核建议。",
        "[0061] 递归检索路径把长问题、候选证据、标准引用和图谱邻域写入受控变量空间，允许模型通过受限函数执行分组、过滤、排序和摘要。允许调用的函数包括search_chunks、get_entity_neighbors、get_standard_clause、group_by_document和summarize_evidence；禁止访问本地任意文件、网络和系统命令。每次函数调用的输入、输出摘要和异常均写入审计历史，最终答案必须引用EvidenceBundle中的证据编号。",
        "[0062] 证据锚定生成器接收EvidenceBundle后，按固定模板输出答案：第一部分为直接结论或保守判断，第二部分为证据依据，第三部分为维修或培训提示，第四部分为需人工复核事项，第五部分为引用来源。对于PACK OFF问题，答案可列出故障现象、可能原因、检查动作、MEL或维修记录要求和培训提示；对于BLEED OFF问题，答案可强调引气系统状态、维修记录和人为因素；对于发动机滑油或飞行操纵问题，答案必须提示安全复核，不把模型输出作为放行依据。",
        "[0063] 图形界面复核包括文档列表、知识树、图谱视图、证据包视图、候选关系队列和审计日志视图。工程人员可查看节点来源页码、关系强度、标准引用和模型置信度；可对候选关系执行确认、驳回、合并、拆分或补充证据；可将复核结论写回经验库。系统对每次人工操作生成AuditRecord，包括操作者、时间、对象、操作前状态、操作后状态和理由。",
        "[0064] 报告输出模块可生成故障分析报告、追溯矩阵、标准引用清单、维修培训提示和审计摘要。报告中的每一条结论均引用EvidenceRecord，不直接引用模型自由文本。若模型服务不可用，系统使用保守模板从检索证据中生成简短回答；若检索索引不可用，系统提示重新构建索引；若图谱存储不可用，系统只输出文档级证据，不输出图谱推理结论。"
    ]),
    ("实施例6：国产计算卡适配和高吞吐token推理部署", [
        "[0065] 本实施例说明硬件适配层和高吞吐token推理链路的可实施方式。硬件适配层不限定单一厂商或单一开源项目，而是把CPU、GPU、NPU、DCU、MLU及其他异构设备抽象为统一DeviceProfile，并根据设备能力选择模型加载、算子后端、张量精度、内存池、并行方式和回退策略。可选设备包括但不限于昇腾NPU、海光DCU、寒武纪MLU、昆仑芯、摩尔线程计算卡以及通用GPU。",
        "[0066] 硬件部署步骤包括：S601，执行设备探测，读取设备数量、驱动版本、算子库版本、可用内存、通信拓扑和健康状态；S602，检查模型格式和后端兼容性，必要时执行格式转换、权重量化或张量并行切分；S603，为嵌入模型、重排序模型和生成模型分别选择后端；S604，创建请求队列、KV cache页池、prefix cache、输出流管理器和审计计数器；S605，发送健康检查请求，确认首token返回、流式输出和日志写入均可用。",
        "[0067] 高吞吐推理请求被拆分为预填充阶段和逐token解码阶段。预填充阶段处理用户问题、检索证据、标准条款和历史经验上下文；逐token解码阶段连续生成答案。请求队列记录request_id、session_id、prompt_tokens、expected_output_tokens、evidence_size、priority、arrival_time和deadline。调度器按输入长度、证据包大小、优先级和设备剩余内存组成动态batch；每个解码步结束后，已完成请求释放缓存，新请求可加入批次，形成连续批处理。",
        "[0068] 分页KV cache采用固定大小页块管理每个会话的键值缓存。CachePage对象包括page_id、request_id、layer_id、device_id、token_start、token_end、ref_count、last_access_time和reusable_flag。若多个请求共享相同标准条款、MEL条目、PACK OFF故障背景或培训规范前缀，系统通过prefix cache复用预填充结果，仅对新增问题和新增证据执行计算。缓存淘汰优先级为已结束会话、低优先级会话、可重建前缀、长时间未访问页和候选问答页。",
        "[0069] 量化和混合精度策略按任务风险和硬件能力选择。培训问答、一般标准检索和资料导航可采用较低比特量化或混合精度；涉及安全评估、DAL建议、放行限制或复杂故障诊断的问题可采用较高精度、较小batch和更严格证据引用。若设备内存不足，系统依次尝试降低batch大小、缩短最大输出长度、启用量化模型、关闭低优先级会话、切换到CPU/GPU或低并发模式，并记录降级原因。",
        "[0070] 面向三份航空维修PDF的并发问答过程如下：第一用户提出PACK OFF灯亮排故问题，系统检索737故障说明中的故障条目、M2文件中的维修记录和培训规范中的MEL训练要求；第二用户提出BLEED OFF与人为因素问题，系统检索引气系统、维修作风和交接复核证据；第三用户提出发动机滑油或飞行操纵风险问题，系统检索故障说明、安全提示和培训规范。调度器把三类请求按证据长度和生成长度组成动态batch，在国产计算卡上执行重排序、摘要和答案生成，并向各用户流式返回。",
        "[0071] 审计统计包括model_version、hardware_backend、device_id、precision_mode、quantization_mode、batch_size、input_tokens、output_tokens、first_token_latency、total_latency、tokens_per_second、cache_hit_rate、kv_pages_used、fallback_event和evidence_ids。上述统计与用户、问题、证据包和答案摘要绑定，形成可追溯日志。对于安全相关问题，系统在答案中保留人工复核提示，并将审计日志用于后续性能调优、硬件容量规划和安全复盘。",
        "[0072] 本实施例的输出为可运行的国产化高吞吐维修问答服务。该服务可在内网或离线环境中部署；当国产计算卡可用时，生成、重排序和证据摘要可由计算卡加速；当计算卡不可用时，检索、证据包整理和保守模板回答仍可在CPU/GPU或低并发模式下运行。由此，本发明在不同单位软硬件条件下均可保持基本可用，并在有加速设备时提升多用户并发场景下的token吞吐。"
    ]),
]


IMPLEMENTATION_FOUR = [
    ("导言", [
        "以下结合附图和实施例，对本发明作进一步详细说明。为避免将同一技术链路拆分为大量零散片段，以下实施例按照工程实现顺序合并为四个大的实施例。每个实施例均给出输入对象、处理步骤、数据结构、判定规则、输出结果和异常回退方式，所属领域技术人员可据此在不依赖特定软件名称或特定硬件厂商的情况下实施本发明。",
    ]),
    ("实施例1：系统部署、基础数据对象和硬件适配初始化", [
        "[0030] 本实施例说明系统从空白部署环境到可接收航空维修资料、可选择推理后端并可写入审计日志的初始化过程。系统至少包括资料接入模块、文档解析模块、语义分块模块、本体知识图谱模块、标准条款检索模块、多角色模型推理模块、证据锚定生成模块、经验库模块、审计日志模块、图形复核模块和硬件适配模块。上述模块可部署在同一工作站、内网服务器或私有化服务集群中，也可按检索服务、图谱服务和生成服务拆分部署。",
        "[0031] 初始化步骤包括：S101，读取配置文件，获得资料目录、输出目录、模型后端、最大上下文长度、检索top_k、稀疏检索权重、稠密检索权重、重排序候选数量、审计日志路径、特征开关和硬件优先级；S102，创建raw、processed、chunk、graph、standard、experience、audit和export目录；S103，加载标准条款库、实体类型表、关系类型表、关联强度权重表、提示词模板和权限策略；S104，执行硬件探测，形成DeviceProfile对象；S105，启动检索索引、图谱存储和模型服务健康检查。",
        "[0032] DeviceProfile对象至少包括device_type、vendor、device_id、memory_total、memory_free、driver_version、backend_name、supported_precision、supported_attention_backend、tensor_parallel_capability、health_status和fallback_priority字段。若未检测到可用加速卡，device_type被设为CPU，系统自动启用低并发、短上下文和保守生成策略；若检测到多个设备，系统按设备健康状态、剩余内存和配置优先级选择生成后端、向量化后端和重排序后端。",
        "[0033] 系统为每个输入文件创建DocumentRecord对象，字段包括document_id、file_name、file_type、source_path、文件指纹、version_tag、security_level、import_time、parser_status、page_count、language、domain_tags和owner。系统的核心中间数据包括ChunkRecord、EntityRecord、RelationRecord、EvidenceRecord、ExperienceEntry、QueryRecord和AuditRecord。上述对象统一通过document_id、source_span、evidence_id和audit_id互相引用，不以本地文件路径作为业务主键。",
        "[0034] ChunkRecord保存chunk_id、document_id、page_no、section_title、chunk_type、text、table_meta、source_span和embedding_status；EntityRecord保存entity_id、name、entity_type、ontology_layer、properties、source_span、standard_refs和confidence；RelationRecord保存relation_id、head_id、tail_id、relation_type、strength、color_level、evidence_ids、standard_refs和audit_status；EvidenceRecord保存evidence_id、document_id、page_no、quote_summary、source_span、standard_refs和reliability；AuditRecord保存operation_id、operator、time、model_version、hardware_backend、parameters、input_summary、output_summary、latency和fallback_event。",
        "[0035] 初始化完成的判定条件为：标准库可查询，资料目录可写入，图谱存储可创建节点和边，检索索引可接受新增chunk，模型后端返回健康状态或启用离线回退模式，审计日志可追加写入。若任一条件失败，系统不进入正式知识抽取流程，而是输出初始化异常报告，列出失败模块、失败原因、回退策略和人工处理建议。"
    ]),
    ("实施例2：资料解析入图、标准注入、多角色协同和六层校验", [
        "[0036] 本实施例说明如何把PDF、Word、Markdown、文本和表格资料转换为可审计知识图谱。资料接入模块接收输入文件后，根据文件类型选择解析器；对于PDF，提取页码、正文、表格、标题层级、页眉页脚和图片说明；对于Word或Markdown，提取段落、标题、表格和列表；对于扫描质量较差的页面，标记为需要OCR或人工复核。解析结果统一写入PageRecord和BlockRecord。",
        "[0037] 语义分块步骤包括：S201，按章节标题、页码、表格边界和段落长度生成初始块；S202，识别块类型，至少包括故障现象块、维修动作块、标准条款块、培训要求块、安全评估块、追溯证据块和一般说明块；S203，对超过长度阈值的块按句子边界继续拆分，对过短且同属一页一节的块进行合并；S204，为每个块生成source_span，记录document_id、page_no、section_title、paragraph_index、char_start和char_end；S205，写入ChunkRecord并进入候选抽取队列。",
        "[0038] 本体约束抽取按照三层本体执行。顶层类型包括实体、过程、功能、质量、角色和倾向；核心层类型包括需求、设计、验证、接口、约束、规范、组件、系统、测试、风险、标准、文档、计划、报告、决策、指标、工具和人员；航空维修领域层类型至少包括故障现象、失效状态、机载系统、部件、维修动作、MEL条目、放行限制、培训科目、人为因素、维修记录、安全目标、DAL分配和证据片段。抽取时，模型或规则引擎只能从允许类型中选择实体类型，不能自由生成未登记类型。",
        "[0039] 实体抽取流程为：S211，从ChunkRecord中识别候选术语；S212，使用术语词典、标准编号规则、英文缩写规则和上下文触发词判断候选类型；S213，将候选实体写成EntityRecord；S214，若同一document_id和相近source_span内出现同名实体，则合并为同一局部实体；S215，若跨文档出现相似实体，则暂存为candidate_same_as关系，等待跨文档共指消解确认。关系抽取流程为：S221，根据关系触发词、句法方向、表格列名和章节语义生成候选关系；S222，根据关系类型允许的头尾实体类型过滤不合法关系；S223，计算关联强度；S224，依据强度区间映射为七级显示等级；S225，将通过阈值的关系写入RelationRecord，将未通过阈值但证据存在的关系写入候选池。",
        "[0040] 标准条款检索采用稀疏检索、稠密检索和类型匹配三路融合。系统从当前ChunkRecord和EntityRecord中抽取故障码、系统名、部件名、MEL编号、标准编号和动作词；BM25路径检索标准编号、中文标题、英文标题和关键词；向量路径检索语义相近条款；类型匹配路径判断条款适用实体类型是否覆盖当前实体类型；最后按R(q,c)=λ1·BM25+λ2·CosSim+λ3·TypeMatch+λ4·ProjectContext计算得分，取Top-K条款作为标准上下文注入抽取和生成流程。",
        "[0041] 多角色协同流程包括知识提取器、本体推理器、安全评估器、追溯检查器、变更影响分析器和文档整理器。知识提取器负责候选实体和关系；本体推理器负责端点类型、层级归属和潜在链接；安全评估器负责危险、失效状态、安全目标、DAL建议和缓解措施；追溯检查器负责需求、设计、测试、维修记录和标准条款之间的覆盖关系；变更影响分析器沿图谱边输出上游、下游、横向、安全和培训影响；文档整理器只对已经通过校验的对象生成证据包和报告。",
        "[0042] 六层校验按固定顺序执行。第一层为格式校验，检查JSON字段、类型、必填项和枚举值；第二层为本体一致性校验，检查实体类型和关系端点是否匹配；第三层为标准引用校验，凡涉及安全目标、DAL、MEL、放行限制、维修记录和培训要求的对象必须绑定标准、规范或证据片段；第四层为证据跨度校验，检查source_span是否能回指到具体文档页码和段落；第五层为冲突检测，比较不同角色或不同文档对同一实体属性、原因关系或处置建议的差异；第六层为置信度阈值校验，低于阈值的对象进入人工复核队列。",
        "[0043] 经验闭环在校验后执行。系统将通过校验的故障分析、维修处置、标准引用、人工复核意见和审计记录组织为ExperienceEntry，字段包括experience_id、title、aircraft_type、task_type、applicable_conditions、content、evidence_ids、standard_refs、confidence、review_status、created_by和usage_count。经验条目先进入临时经验库，经工程人员审核后可提升为永久经验。后续遇到相同机型、相同系统、相似故障或相同标准条款时，系统把匹配经验作为上下文注入，但经验只作为辅助约束，不替代标准条款和原文证据。"
    ]),
    ("实施例3：围绕三份航空维修PDF的故障知识图谱、问答和复核输出", [
        "[0044] 本实施例以三份实际航空维修资料为输入，说明从PDF资料到故障知识图谱、故障树、因果树、跨文档证据包和维修问答的完整实施过程。输入资料包括《737典型故障说明R1（飞行版）.pdf》《M2-航空器维修（第一次修订）R1版2025年4月7日.pdf》和《航空器维修基础知识和实作培训规范-正式版.pdf》。三份资料分别提供故障案例、维修管理和培训实作规范，系统将其作为同一航空维修知识域的不同证据源处理。",
        "[0045] 对《737典型故障说明R1（飞行版）.pdf》，系统执行以下步骤：S301，解析故障目录和故障详情页，识别PACK OFF、BLEED OFF、DUAL BLEED OFF、发动机滑油、飞行操纵和起飞构型等故障主题；S302，为每个故障主题建立Episode对象；S303，在Episode下建立故障现象、运行条件、可能原因、排故动作、排除原因、维修建议和预防提示等Facet；S304，把每个Facet中的具体表述转化为FacetPoint并关联EntityRecord；S305，依据触发词和章节语义建立symptom_of、caused_by、requires_check、excluded_cause和maintenance_action等关系。",
        "[0046] 以PACK OFF为例，系统把“PACK OFF灯亮”“引气压力异常”“空调组件工作异常”“查阅MEL”“完成维修记录”分别抽取为故障现象、失效状态、部件状态、标准条款和维修动作实体。若同一段落同时出现“压力低”和“组件关断”，系统不直接将二者互为因果，而是检查上下文是否存在“导致、引起、由于、检查、排除”等触发词，并结合MEL或维修程序引用确定关系方向。最终输出的PACK OFF子图至少包含故障现象节点、引气系统节点、空调组件节点、MEL条目节点、维修动作节点和培训提示节点。",
        "[0047] 对《M2-航空器维修（第一次修订）R1版2025年4月7日.pdf》，系统重点抽取维修作风、人为因素、维修记录、放行要求、工具管理、工卡执行和诚信管理等实体；对《航空器维修基础知识和实作培训规范-正式版.pdf》，系统重点抽取培训科目、实作步骤、考核要点、安全注意事项和MEL查阅训练要求。维修文件中的record_requirement、human_factor和procedure_requirement与故障说明中的maintenance_action关联；培训规范中的training_item和assessment_requirement与故障案例中的故障主题和维修动作关联。",
        "[0048] 跨文档关联分为共指消解、标准桥接和图结构确认三步。共指消解统一“空调组件”“PACK组件”“组件关断”等近义表达；标准桥接把MEL条目、维修记录要求和培训规范条款连接到相同或相关的标准对象；图结构确认检查两个实体是否共享部件、系统、动作或标准引用。只有同时满足名称相似、类型兼容和证据支持的关联才写入正式图谱，其他关联保留为候选关系。",
        "[0049] 故障树生成以顶事件为入口。系统接收顶事件名称，例如“引气系统压力低”或“PACK OFF灯亮”，在图谱中定位对应failure_state节点；沿caused_by、failure_leads_to和requires_check关系向上展开原因节点；将可直接检查或可维修的部件、传感器、活门、管路和控制逻辑标记为基本事件；根据多个原因是否独立、共同作用或排除关系推断OR门、AND门或NOT标记。因果树生成以时间和证据强度为主线，从故障现象开始，沿confirmed优先、possible次之、candidate再次之的顺序组织因果链。",
        "[0050] 用户问答时，系统首先创建QueryRecord，字段包括query_id、user_id、query_text、query_type、aircraft_type、priority、created_time、route、status和audit_id。若问题包含明确故障码、标准编号、部件名或MEL条目，则优先使用基线混合检索路径；若问题包含“为什么、如何排故、影响哪些、关联哪些标准”等多步推理表达，则使用Agentic多步路径；若问题要求综合三份资料或列出所有证据，则使用递归检索路径；若高级路径未启用或模型不可用，则降级到基线混合检索路径。",
        "[0051] 证据锚定生成器接收EvidenceBundle后，按固定模板输出答案：第一部分为直接结论或保守判断，第二部分为证据依据，第三部分为维修或培训提示，第四部分为需人工复核事项，第五部分为引用来源。对于PACK OFF问题，答案可列出故障现象、可能原因、检查动作、MEL或维修记录要求和培训提示；对于BLEED OFF问题，答案可强调引气系统状态、维修记录和人为因素；对于发动机滑油或飞行操纵问题，答案必须提示安全复核，不把模型输出作为放行依据。",
        "[0052] 图形界面复核包括文档列表、知识树、图谱视图、证据包视图、候选关系队列和审计日志视图。工程人员可查看节点来源页码、关系强度、标准引用和模型置信度；可对候选关系执行确认、驳回、合并、拆分或补充证据；可将复核结论写回经验库。系统对每次人工操作生成AuditRecord，包括操作者、时间、对象、操作前状态、操作后状态和理由。报告输出模块可生成故障分析报告、追溯矩阵、标准引用清单、维修培训提示和审计摘要，报告中的每一条结论均引用EvidenceRecord。"
    ]),
    ("实施例4：国产计算卡适配和高吞吐token推理部署", [
        "[0053] 本实施例说明硬件适配层和高吞吐token推理链路的可实施方式。硬件适配层不限定单一厂商或单一开源项目，而是把CPU、GPU、NPU、DCU、MLU及其他异构设备抽象为统一DeviceProfile，并根据设备能力选择模型加载、算子后端、张量精度、内存池、并行方式和回退策略。可选设备包括但不限于昇腾NPU、海光DCU、寒武纪MLU、昆仑芯、摩尔线程计算卡以及通用GPU。",
        "[0054] 硬件部署步骤包括：S401，执行设备探测，读取设备数量、驱动版本、算子库版本、可用内存、通信拓扑和健康状态；S402，检查模型格式和后端兼容性，必要时执行格式转换、权重量化或张量并行切分；S403，为嵌入模型、重排序模型和生成模型分别选择后端；S404，创建请求队列、KV cache页池、prefix cache、输出流管理器和审计计数器；S405，发送健康检查请求，确认首token返回、流式输出和日志写入均可用。",
        "[0055] 高吞吐推理请求被拆分为预填充阶段和逐token解码阶段。预填充阶段处理用户问题、检索证据、标准条款和历史经验上下文；逐token解码阶段连续生成答案。请求队列记录request_id、session_id、prompt_tokens、expected_output_tokens、evidence_size、priority、arrival_time和deadline。调度器按输入长度、证据包大小、优先级和设备剩余内存组成动态batch；每个解码步结束后，已完成请求释放缓存，新请求可加入批次，形成连续批处理。",
        "[0056] 分页KV cache采用固定大小页块管理每个会话的键值缓存。CachePage对象包括page_id、request_id、layer_id、device_id、token_start、token_end、ref_count、last_access_time和reusable_flag。若多个请求共享相同标准条款、MEL条目、PACK OFF故障背景或培训规范前缀，系统通过prefix cache复用预填充结果，仅对新增问题和新增证据执行计算。缓存淘汰优先级为已结束会话、低优先级会话、可重建前缀、长时间未访问页和候选问答页。",
        "[0057] 量化和混合精度策略按任务风险和硬件能力选择。培训问答、一般标准检索和资料导航可采用较低比特量化或混合精度；涉及安全评估、DAL建议、放行限制或复杂故障诊断的问题可采用较高精度、较小batch和更严格证据引用。若设备内存不足，系统依次尝试降低batch大小、缩短最大输出长度、启用量化模型、关闭低优先级会话、切换到CPU/GPU或低并发模式，并记录降级原因。",
        "[0058] 面向三份航空维修PDF的并发问答过程如下：第一用户提出PACK OFF灯亮排故问题，系统检索737故障说明中的故障条目、M2文件中的维修记录和培训规范中的MEL训练要求；第二用户提出BLEED OFF与人为因素问题，系统检索引气系统、维修作风和交接复核证据；第三用户提出发动机滑油或飞行操纵风险问题，系统检索故障说明、安全提示和培训规范。调度器把三类请求按证据长度和生成长度组成动态batch，在国产计算卡上执行重排序、摘要和答案生成，并向各用户流式返回。",
        "[0059] 审计统计包括model_version、hardware_backend、device_id、precision_mode、quantization_mode、batch_size、input_tokens、output_tokens、first_token_latency、total_latency、tokens_per_second、cache_hit_rate、kv_pages_used、fallback_event和evidence_ids。上述统计与用户、问题、证据包和答案摘要绑定，形成可追溯日志。对于安全相关问题，系统在答案中保留人工复核提示，并将审计日志用于后续性能调优、硬件容量规划和安全复盘。",
        "[0060] 本实施例的输出为可运行的国产化高吞吐维修问答服务。该服务可在内网或离线环境中部署；当国产计算卡可用时，生成、重排序和证据摘要可由计算卡加速；当计算卡不可用时，检索、证据包整理和保守模板回答仍可在CPU/GPU或低并发模式下运行。由此，本发明在不同单位软硬件条件下均可保持基本可用，并在有加速设备时提升多用户并发场景下的token吞吐。"
    ]),
]


def rewrite_fourth_section():
    shutil.copy2(DOCX, BACKUP)
    doc = Document(DOCX)
    start = find_para(doc, "第四部分：具体实施方式")
    end = find_para(doc, "第五部分：发明效果")
    # Keep the fourth title, remove old implementation paragraphs.
    for p in list(doc.paragraphs[start + 1:end]):
        remove_para(p)
    anchor = doc.paragraphs[start]
    flat = []
    for title, paras in IMPLEMENTATION_FOUR:
        if title == "导言":
            flat.extend((p, False) for p in paras)
        else:
            flat.append((title, True))
            flat.extend((p, False) for p in paras)
    for content, bold in reversed(flat):
        add_after(anchor, content, bold=bold)
    doc.save(DOCX)


def update_report():
    note = f"""

## 第四部分四实施例重写

补充时间：{datetime.now():%Y-%m-%d %H:%M:%S}

根据进一步复核意见，已将第四部分原67个短实施例合并重写为4个大的、可按步骤实施的实施例：系统部署、基础数据对象和硬件适配初始化；资料解析入图、标准注入、多角色协同和六层校验；围绕三份航空维修PDF的故障知识图谱、问答和复核输出；国产计算卡适配和高吞吐token推理部署。每个实施例均补充了输入对象、步骤编号、关键字段、判定条件、输出结果和异常回退方式。
"""
    REPORT.write_text(REPORT.read_text(encoding="utf-8") + note, encoding="utf-8")


def checks():
    doc = Document(DOCX)
    txt = "\n".join(p.text for p in doc.paragraphs)
    bad = ["综合实施例", "软件设计说明书", "第九部分：软件设计说明书支撑摘录", "待补充", "TODO", "占位图", "示意图占位"]
    with zipfile.ZipFile(DOCX) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    return {
        "bad_terms": [b for b in bad if b in txt],
        "embodiments": sum(1 for p in doc.paragraphs if p.text.strip().startswith("实施例")),
        "paragraphs": len(doc.paragraphs),
        "media": len(media),
    }


if __name__ == "__main__":
    rewrite_fourth_section()
    update_report()
    print(f"backup {BACKUP}")
    print(checks())
