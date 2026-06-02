#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成发明专利和软件著作权申请Word文档
"""
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
except ImportError:
    print("正在安装python-docx...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE

def create_patent_document():
    """创建发明专利申请文档"""
    doc = Document()
    
    # 设置标题
    title = doc.add_heading('发明专利申请文件', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 第一部分：发明专利请求书
    doc.add_heading('第一部分：发明专利请求书', level=1)
    
    doc.add_heading('一、发明名称', level=2)
    doc.add_paragraph('一种基于工业本体和大语言模型的民机适航知识图谱智能构建与分析方法')
    
    doc.add_heading('二、发明人信息', level=2)
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    cells = table.rows[0].cells
    cells[0].text = '项目'
    cells[1].text = '内容'
    table.rows[1].cells[0].text = '发明人'
    table.rows[1].cells[1].text = '（请填写全部发明人姓名）'
    table.rows[2].cells[0].text = '第一发明人'
    table.rows[2].cells[1].text = '（请填写）'
    table.rows[3].cells[0].text = '国籍'
    table.rows[3].cells[1].text = '中国'
    
    doc.add_heading('三、申请人信息', level=2)
    table = doc.add_table(rows=6, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = '项目'
    table.rows[0].cells[1].text = '内容'
    table.rows[1].cells[0].text = '申请人'
    table.rows[1].cells[1].text = '（请填写单位/个人全称）'
    table.rows[2].cells[0].text = '类型'
    table.rows[2].cells[1].text = '（职务发明/非职务发明）'
    table.rows[3].cells[0].text = '地址'
    table.rows[3].cells[1].text = '（请填写）'
    table.rows[4].cells[0].text = '邮编'
    table.rows[4].cells[1].text = '（请填写）'
    table.rows[5].cells[0].text = '联系人'
    table.rows[5].cells[1].text = '（请填写）'
    
    doc.add_heading('四、分类号（建议）', level=2)
    doc.add_paragraph('IPC分类号：G06F 16/36（本体/知识图谱）; G06F 40/30（自然语言处理）; G06N 5/04（知识处理系统）; B64F 5/00（航空器维修/检查辅助装置）')
    doc.add_paragraph('CPC分类号：G06F 16/367; G06N 20/00')
    
    doc.add_heading('五、摘要', level=2)
    abstract = '''本发明公开了一种基于工业本体和大语言模型的民机适航知识图谱智能构建与分析方法。该方法采用BFO/IOF三层本体架构（基本形式本体顶层→工业本体基础核心层→航空领域层），定义了74种实体类型和36种关系类型，并引入赤橙黄绿青蓝紫七级关联强度色谱量化知识关联的紧密程度。方法通过多角色大语言模型Agent（包括知识提取器、本体推理器、安全评估师、V模型追溯器、变更影响分析器和文档整理器六个专业角色），从多格式民机技术文档中自动提取结构化知识并构建知识图谱。系统集成中、美、欧三大适航当局标准体系作为知识参考源，并具备分析过程中自动提炼和积累工程经验的能力，形成"分析→提炼→积累→注入"的经验闭环。本发明支持本地离线运行和国产化华为昇腾NPU部署，适用于民用航空器适航认证的安全评估、追溯分析和变更影响分析等工程场景。'''
    doc.add_paragraph(abstract)
    
    doc.add_paragraph('关键词：工业本体；知识图谱；大语言模型；适航认证；安全评估；关联强度色谱；经验积累')
    
    # 第二部分：权利要求书
    doc.add_page_break()
    doc.add_heading('第二部分：权利要求书', level=1)
    
    doc.add_heading('权利要求1（独立权利要求——方法）', level=2)
    claim1 = '''一种基于工业本体和大语言模型的民机适航知识图谱智能构建与分析方法，其特征在于，包括以下步骤：

S1. 构建三层本体架构：建立基于BFO基本形式本体的顶层本体、基于IOF工业本体基础的核心层本体和面向民机适航领域的领域层本体，所述三层本体定义实体类型层级和关系类型集合，其中实体类型层级包含BFO顶层实体类型、IOF核心层实体类型和航空领域层实体类型，关系类型集合包含追溯关系、安全评估关系、认证关系和变更管理关系；

S2. 定义关联强度色谱：将知识关系的关联强度量化为0至1的连续数值，并映射到赤、橙、黄、绿、青、蓝、紫七级色谱区间，其中赤色表示关联强度0.86至1.00的极强关联，紫色表示关联强度0.00至0.15的潜在关联，中间色谱依次对应递减的关联强度区间；

S3. 多角色大语言模型知识提取：配置至少包含知识提取角色和本体推理角色的多角色大语言模型Agent系统，其中每个角色具有独立的系统提示词，所述系统提示词包含步骤S1所述三层本体的类型体系和步骤S2所述的关联强度标注规则；所述知识提取角色从输入的民机技术文档中提取结构化的实体和关系信息，所述本体推理角色基于已构建的知识图谱执行深度推理分析；

S4. 知识图谱构建：将步骤S3提取的实体作为图谱节点，关系作为图谱带权边存入RDF三元组存储和图数据结构中，每条边携带步骤S2所述的关联强度值，形成可查询、可推理的知识图谱；

S5. 适航标准注入：建立包含中国CAAC、美国FAA和欧洲EASA三大适航当局标准数据库，在步骤S3的知识提取过程中，根据文档关键词自动检索相关标准条款，将标准条款摘要注入大语言模型的提示词上下文，使知识提取结果符合适航标准要求；

S6. 经验闭环积累：在步骤S3完成后，自动调用大语言模型对分析结果进行经验提炼，将提炼的经验按永久经验和临时经验分类存储，并在后续分析中将历史经验注入大语言模型的提示词上下文，形成"分析→提炼→积累→注入"的经验闭环。'''
    doc.add_paragraph(claim1)
    
    # 继续添加其他权利要求...
    for i in range(2, 11):
        doc.add_heading(f'权利要求{i}', level=2)
        if i == 2:
            doc.add_paragraph('根据权利要求1所述的方法，其特征在于，步骤S1中所述的实体类型层级具体包括：\n（a）BFO顶层的6种基础类型：实体、过程、功能、质量、角色、倾向；\n（b）IOF核心层的18种工业通用类型：需求、设计、验证、接口、约束、规范、组件、系统、测试、风险、标准、文档、计划、报告、决策、指标、工具、人员；\n（c）航空领域层的50种民机适航专用类型。')
        elif i == 8:
            doc.add_paragraph('一种基于工业本体和大语言模型的民机适航知识图谱智能构建与分析系统，其特征在于，包括：本体知识图谱引擎模块、大语言模型多角色Agent引擎模块、适航标准库模块、经验积累与审计引擎模块、多格式文档解析模块、图形用户界面模块。')
    
    # 保存文档
    output_path = os.path.join(os.path.dirname(__file__), '发明专利申请文件.docx')
    doc.save(output_path)
    print(f"发明专利申请文档已生成: {output_path}")
    return output_path

def create_software_document():
    """创建软件著作权申请文档"""
    doc = Document()
    
    # 设置标题
    title = doc.add_heading('计算机软件著作权登记申请文件', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 第一部分：软件著作权登记申请表
    doc.add_heading('第一部分：软件著作权登记申请表', level=1)
    
    doc.add_heading('一、软件基本信息', level=2)
    table = doc.add_table(rows=8, cols=2)
    table.style = 'Table Grid'
    data = [
        ('项目', '内容'),
        ('软件全称', '基于工业本体的民机适航知识图谱智能分析系统'),
        ('软件简称', '民机适航知识图谱系统'),
        ('版本号', 'V4.0'),
        ('软件分类', '应用软件'),
        ('开发完成日期', '2026年04月16日'),
        ('首次发表日期', '2026年04月16日'),
        ('开发方式', '独立开发'),
    ]
    for i, (k, v) in enumerate(data):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
    
    doc.add_heading('二、软件开发运行环境', level=2)
    doc.add_heading('2.1 开发环境', level=3)
    table = doc.add_table(rows=6, cols=2)
    table.style = 'Table Grid'
    data = [
        ('项目', '内容'),
        ('开发语言', 'Python 3.10+'),
        ('开发工具', 'Visual Studio Code / PyCharm'),
        ('操作系统', 'Windows 10/11、Linux (Ubuntu 22.04+)'),
        ('数据库', 'JSON文件存储、RDF三元组存储'),
        ('AI推理框架', 'Ollama本地推理 / vLLM-Ascend（华为昇腾910B NPU）'),
    ]
    for i, (k, v) in enumerate(data):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
    
    doc.add_heading('三、软件功能说明', level=2)
    doc.add_paragraph('本软件是面向民用航空器适航认证领域的智能知识图谱分析系统，基于工业本体基础框架（IOF/BFO三层架构），集成大语言模型推理能力，为空客、波音等民机设计团队提供从需求分析、安全评估到认证追溯的全流程智能化工具。')
    
    # 功能模块列表
    modules = [
        ('模块一：通用本体知识图谱引擎', '基于BFO/IOF三层架构，支持74种实体类型、36种关系类型、七级关联强度色谱，提供RDF三元组存储和NetworkX图分析双引擎。'),
        ('模块二：大语言模型多角色Agent引擎', 'Ollama/vLLM-Ascend双后端架构，六大Agent角色（知识提取器、本体推理器、文档整理器、安全评估师、V模型追溯器、变更影响分析器）。'),
        ('模块三：适航认证工作流引擎', 'ARP4761安全评估、DO-178C追溯完整性检查、ARP4754A V模型全生命周期、变更影响分析。'),
        ('模块四：适航标准库', '中、美、欧三大适航当局23项核心标准，交叉引用对照，自动上下文注入。'),
        ('模块五：经验积累与审计引擎', '自动经验提炼、按机型分目录管理、中英文双语文档生成、人工审核流程。'),
        ('模块六：多格式文档解析', '支持PDF/DOCX/XLSX/PPTX/MD等14种格式。'),
        ('模块七：可视化与交互界面', '暗色航空主题、知识图谱可视化、七大功能标签页。'),
    ]
    
    for name, desc in modules:
        doc.add_heading(name, level=3)
        doc.add_paragraph(desc)
    
    doc.add_heading('四、软件技术特点', level=2)
    features = [
        '工业本体驱动：采用IOF/BFO国际工业本体框架，实现航空领域知识的专业化、标准化建模。',
        '国产化适配：支持华为昇腾910B NPU部署，通过vLLM-Ascend框架实现国产AI芯片推理加速。',
        '全离线运行：系统无需互联网连接，满足军工/航空航天保密环境要求。',
        '三方标准融合：同时覆盖CAAC/FAA/EASA标准体系，适用于双边/多边适航认证项目。',
        '经验闭环：设计"分析→提炼→积累→注入"的经验闭环，使Agent随使用不断增强。',
    ]
    for feature in features:
        doc.add_paragraph(f'• {feature}')
    
    # 保存文档
    output_path = os.path.join(os.path.dirname(__file__), '软件著作权申请文件.docx')
    doc.save(output_path)
    print(f"软件著作权申请文档已生成: {output_path}")
    return output_path

if __name__ == '__main__':
    base_dir = r'C:\Users\weiku\Desktop\知识库\知识产权'
    os.chdir(base_dir)
    
    print("开始生成Word文档...")
    patent_path = create_patent_document()
    software_path = create_software_document()
    print("文档生成完成！")
