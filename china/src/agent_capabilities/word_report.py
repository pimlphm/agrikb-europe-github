from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any


def build_word_report(
    output_path: str | Path,
    *,
    title: str,
    query: str,
    generation: dict[str, Any],
    tree: dict[str, Any],
    capabilities: dict[str, Any],
    graph_image_path: str | Path | None = None,
    graph_svg_path: str | Path | None = None,
    graph_html_path: str | Path | None = None,
) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as exc:  # pragma: no cover - depends on runtime install
        raise RuntimeError("python-docx is required for Word report export.") from exc

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)

    heading = document.add_heading(title, 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Generated at {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    document.add_heading("1. Query", level=1)
    document.add_paragraph(query or "Current knowledge-base overview")

    document.add_heading("2. Executive Answer", level=1)
    for line in _answer_lines(generation.get("answer", "")):
        if line.startswith("###"):
            document.add_heading(line.lstrip("# ").strip(), level=2)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)

    document.add_heading("3. Evidence Chain", level=1)
    evidence = _evidence_items(generation)
    if evidence:
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["ID", "Source", "Support", "Excerpt"]
        for idx, label in enumerate(headers):
            table.rows[0].cells[idx].text = label
        for item in evidence[:12]:
            row = table.add_row().cells
            row[0].text = str(item.get("id") or item.get("chunk_id") or "")
            row[1].text = str(item.get("doc_name") or item.get("source") or "")
            row[2].text = str(item.get("supports") or item.get("chunk_type") or "")
            row[3].text = _compact(item.get("excerpt") or item.get("text") or "", 320)
    else:
        document.add_paragraph("No evidence chain was returned for this run.")

    document.add_heading("4. Knowledge Graph Snapshot", level=1)
    document.add_paragraph(
        f"Documents: {tree.get('document_count', 0)}; chunks: {tree.get('chunk_count', 0)}. "
        "The graph export uses a no-overlap hierarchical layout and a separate dynamic HTML view for pan/zoom/expand operations."
    )
    if graph_image_path and Path(graph_image_path).exists():
        try:
            document.add_picture(str(graph_image_path), width=Inches(6.7))
        except Exception:
            document.add_paragraph(f"Graph PNG: {graph_image_path}")
    if graph_svg_path:
        document.add_paragraph(f"High-resolution SVG: {graph_svg_path}")
    if graph_html_path:
        document.add_paragraph(f"Dynamic HTML graph: {graph_html_path}")

    document.add_heading("5. Rowboat-Style Agent Capabilities", level=1)
    document.add_paragraph(
        "This backend integrates Rowboat-style capabilities as a local adapter: durable Markdown memory, "
        "inspectable graph context, reviewable workflows, and exportable artifacts."
    )
    for item in capabilities.get("capabilities", []):
        document.add_paragraph(f"{item.get('label')} ({item.get('status')}): {item.get('endpoint')}", style="List Bullet")

    document.add_heading("6. Residual Risk And Review Notes", level=1)
    document.add_paragraph(
        "Answers are evidence-grounded but still require domain expert review before maintenance release, "
        "airworthiness compliance decisions, patent filing, or formal engineering sign-off."
    )

    document.save(output)
    return output


def _answer_lines(answer: str) -> list[str]:
    lines = [line.strip() for line in str(answer or "").splitlines()]
    return [line for line in lines if line]


def _evidence_items(generation: dict[str, Any]) -> list[dict[str, Any]]:
    if generation.get("evidence_chain"):
        return list(generation["evidence_chain"])
    chunks = generation.get("evidence_chunks") or generation.get("retrieval_results") or []
    items = []
    for index, chunk in enumerate(chunks[:12], start=1):
        items.append(
            {
                "id": f"E{index}",
                "doc_name": chunk.get("doc_name") or chunk.get("source", ""),
                "supports": chunk.get("chunk_type", "evidence"),
                "excerpt": chunk.get("text", ""),
                "score": chunk.get("score", ""),
            }
        )
    return items


def _compact(value: str, limit: int) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "..."
