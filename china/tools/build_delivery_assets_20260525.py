# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import shutil
import subprocess
from datetime import date
from pathlib import Path

import imageio_ffmpeg
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
DELIVERABLES = ROOT / "deliverables" / "AgriKB_DualTrack_FirstPrize_Materials_20260523"
SCREENSHOT_DIR = ROOT / "artifacts" / "screenshots" / "latest_packaging_20260525"
SCREENSHOT_EXPORT = DELIVERABLES / "screenshots_20260525"
VIDEO_PATH = DELIVERABLES / "AgriKB_function_demo_recording_20260525.mp4"
PPTX_PATH = DELIVERABLES / "AgriKB_数智新农人经营中枢_路演终版_20260525.pptx"
QR_PATH = DELIVERABLES / "AgriKB_mobile_test_qr.png"
TODAY = date(2026, 5, 25)


GREEN = "1F6B45"
DEEP = "153B2B"
GOLD = "B8842E"
RED = "B34536"
INK = "1C251F"
MUTED = "65746A"
PAPER = "F6F8F1"
LINE = "D7E3D8"
WHITE = "FFFFFF"


SCREENSHOT_META = [
    ("01_home_clean_v90.png", "首页：一线入口更少、更清楚", "清空历史对话后，农户、农业局和企业从同一套入口进入。"),
    ("02_greenhouse_monitoring_v90.png", "大棚监测：指标、动画和建议同屏", "温湿度、光照、土壤水分和 EC 形成可执行的棚室管理提醒。"),
    ("03_greenhouse_detail_v90.png", "每个按钮都能继续点开", "点开棚温、湿度、材料、窗口等节点后，会出现更细的解释和下一步分析入口。"),
    ("04_information_aggregation_v90.png", "资讯聚合：政策、天气、行情、收购同步筛选", "像头条信息流一样聚合，但以农业经营判断为核心。"),
    ("05_competition_headlines_no_badges_v90.png", "创业大赛专题：六类热点聚合", "右上角数量角标已去掉，展开后仍保留每类 10 条实时头条。"),
    ("06_market_price_loop_v90.png", "行情与农资：价格循环滚动并给出建议", "用户查品名后，品类信息沉淀到本体文件夹。"),
    ("07_weather_horizons_v90.png", "农业气象：今天到三周趋势滑页", "默认江苏句容，也允许切换城市和地图位置。"),
    ("08_policy_drilldown_v90.png", "政策雷达：对象、材料、窗口、部门可点开", "政策解读、申报准备和本地部门咨询聚合到一张操作图。"),
]


def hex_to_rgb(value: str) -> tuple[int, int, int]:
    value = value.strip("#")
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


def ppt_color(value: str) -> PptRGB:
    r, g, b = hex_to_rgb(value)
    return PptRGB(r, g, b)


def ensure_dirs() -> None:
    DELIVERABLES.mkdir(parents=True, exist_ok=True)
    SCREENSHOT_EXPORT.mkdir(parents=True, exist_ok=True)


def copy_latest_screenshots() -> list[Path]:
    copied: list[Path] = []
    for filename, _, _ in SCREENSHOT_META:
        src = SCREENSHOT_DIR / filename
        if src.exists():
            dst = SCREENSHOT_EXPORT / filename
            shutil.copy2(src, dst)
            copied.append(dst)
    if not copied:
        raise FileNotFoundError(f"No screenshots found in {SCREENSHOT_DIR}")
    return copied


def load_font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc") if bold else Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if draw.textlength(candidate, font=font) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def make_demo_video(screenshots: list[Path]) -> Path:
    frame_dir = DELIVERABLES / "_video_frames_20260525"
    if frame_dir.exists():
        shutil.rmtree(frame_dir)
    frame_dir.mkdir(parents=True, exist_ok=True)

    title_font = load_font(34, True)
    note_font = load_font(18)
    small_font = load_font(14, True)
    index_font = load_font(16)
    bg = hex_to_rgb(PAPER)
    deep = hex_to_rgb(DEEP)
    green = hex_to_rgb(GREEN)
    muted = hex_to_rgb(MUTED)
    white = hex_to_rgb(WHITE)
    line = hex_to_rgb(LINE)

    frame_index = 0
    meta_by_name = {name: (title, note) for name, title, note in SCREENSHOT_META}
    for shot in screenshots:
        title, note = meta_by_name.get(shot.name, ("AgriKB 功能演示", "最新界面截图"))
        with Image.open(shot) as raw:
            img = raw.convert("RGB")
        canvas = Image.new("RGB", (1280, 720), bg)
        draw = ImageDraw.Draw(canvas)
        draw.rectangle((54, 38, 108, 42), fill=green)
        draw.text((118, 28), "AgriKB 路演功能讲解 2026-05-25", fill=green, font=small_font)
        draw.text((54, 74), title, fill=deep, font=title_font)
        y = 122
        for line_text in wrap_text(draw, note, note_font, 780):
            draw.text((58, y), line_text, fill=muted, font=note_font)
            y += 26
        draw.rounded_rectangle((70, 168, 1210, 646), radius=18, fill=white, outline=line, width=2)
        max_w, max_h = 1100, 438
        scale = min(max_w / img.width, max_h / img.height)
        new_size = (int(img.width * scale), int(img.height * scale))
        resized = img.resize(new_size, Image.Resampling.LANCZOS)
        x = 90 + (1100 - new_size[0]) // 2
        y_img = 188 + (438 - new_size[1]) // 2
        canvas.paste(resized, (x, y_img))
        draw.rounded_rectangle((70, 664, 1210, 694), radius=15, fill=hex_to_rgb("EAF4EA"), outline=hex_to_rgb("C9DCCB"), width=1)
        draw.text((104, 670), "演示路径：点入口 → 看详情 → 新农人助手聚合政策、行情、天气、路线与证据 → 形成可执行建议", fill=deep, font=small_font)
        draw.text((1160, 34), f"{screenshots.index(shot) + 1}/{len(screenshots)}", fill=muted, font=index_font)
        for _ in range(5):
            frame_index += 1
            canvas.save(frame_dir / f"frame_{frame_index:04d}.png")

    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    if VIDEO_PATH.exists():
        VIDEO_PATH.unlink()
    cmd = [
        ffmpeg,
        "-y",
        "-framerate",
        "2",
        "-i",
        str(frame_dir / "frame_%04d.png"),
        "-c:v",
        "libx264",
        "-pix_fmt",
        "yuv420p",
        "-movflags",
        "+faststart",
        str(VIDEO_PATH),
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    shutil.rmtree(frame_dir)
    return VIDEO_PATH


def add_doc_run(paragraph, text: str, bold: bool = False, color: str = INK, size: float | None = None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def set_doc_style(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)
    for name, size, color in [("Heading 1", 17, GREEN), ("Heading 2", 13, DEEP), ("Heading 3", 11.5, GOLD)]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    doc.core_properties.title = title
    doc.core_properties.subject = "AgriKB 数智新农人经营中枢"
    doc.core_properties.author = "AgriKB Project"
    doc.core_properties.comments = "2026-05-25 one-click trusted PC delivery refresh."


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_cover(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    add_doc_run(p, "AgriKB", True, GREEN, 13)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    add_doc_run(p, title, True, DEEP, 24)
    p = doc.add_paragraph()
    add_doc_run(p, subtitle, False, MUTED, 12)
    table = doc.add_table(rows=4, cols=2)
    rows = [
        ("材料版本", "V1.1 交付更新版"),
        ("更新日期", TODAY.isoformat()),
        ("产品方向", "数智信息 × 新农人 × 县域农业经营中枢"),
        ("部署形态", "可信 Windows 新电脑解压后一键启动"),
    ]
    for row, (k, v) in zip(table.rows, rows):
        row.cells[0].text = k
        row.cells[1].text = v
        shade_cell(row.cells[0], "EAF4EA")
    doc.add_paragraph()
    p = doc.add_paragraph()
    add_doc_run(p, "说明：", True, GREEN)
    add_doc_run(p, "本文档随路演 PPT、演示视频、截图、专利/软著初稿和便携运行包同步更新。API key 仅写入本地 .env 配置，不在文档正文展示。", False, INK)
    doc.add_page_break()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        add_doc_run(p, item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        add_doc_run(p, item)


def build_docx(filename: str, title: str, subtitle: str, sections: list[tuple[str, list[str], list[str] | None]]) -> Path:
    doc = Document()
    set_doc_style(doc, title)
    add_cover(doc, title, subtitle)
    for heading, paras, bullets in sections:
        doc.add_heading(heading, level=1)
        for para in paras:
            p = doc.add_paragraph()
            add_doc_run(p, para)
        if bullets:
            add_bullets(doc, bullets)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_doc_run(footer, f"AgriKB 数智新农人经营中枢 | V1.1 | {TODAY.isoformat()}", False, MUTED, 8.5)
    path = DELIVERABLES / filename
    doc.save(path)
    return path


def write_md(filename: str, title: str, sections: list[tuple[str, list[str], list[str] | None]]) -> Path:
    lines = [f"# {title}", "", f"- 版本：V1.1 交付更新版", f"- 日期：{TODAY.isoformat()}", ""]
    for heading, paras, bullets in sections:
        lines.extend([f"## {heading}", ""])
        for para in paras:
            lines.extend([para, ""])
        if bullets:
            for item in bullets:
                lines.append(f"- {item}")
            lines.append("")
    path = DELIVERABLES / filename
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def build_documents() -> list[Path]:
    doc_specs = [
        (
            "01_AgriKB_使用说明书.docx",
            "01_AgriKB_使用说明书.md",
            "AgriKB 使用说明书",
            "面向新农人、农业局和农业企业的一线操作指南",
            [
                ("一键启动", ["在可信 Windows 新电脑上解压交付包后，双击“ 一键启动AgriKB.bat ”即可启动本地服务并打开浏览器。默认地址为 http://127.0.0.1:8010/ui/。"], ["首次启动会使用包内便携 Python 与已安装依赖，不要求新电脑预装 Python。", "如端口被占用，可运行 Start-TrustedPC.ps1 -Port 其他端口。"]),
                ("日常使用", ["用户从首页四个经营入口或左侧功能入口进入：政策、行情、气象、质量溯源、大棚监测和资讯聚合。"], ["点击任意模块卡片后，可继续点开里面的指标、节点和流程按钮。", "需要综合判断时，点击“让新农人助手继续分析”，系统会自动聚合当前模块、天气、行情、路线、政策和头条趋势。", "回答下方的小喇叭可朗读；语音输入功能已移除。"]),
                ("文件与设置", ["右上角“文件”用于资料导入、导出和知识沉淀；“设置”用于修改 API key、模型、联网开关和显示选项。"], ["联网搜索可开关控制。", "API key 写在本地 .env 文件，不应上传到公开仓库。"]),
            ],
        ),
        (
            "02_AgriKB_产品说明书.docx",
            "02_AgriKB_产品说明书.md",
            "AgriKB 产品说明书",
            "数智信息 × 新农人融合项目说明",
            [
                ("产品定位", ["AgriKB 是面向县域农业场景的数智新农人经营中枢，把政策、市场、气象、路线、质量溯源、作物监测和农业知识库合到一个自然语言入口里。"], ["服务对象包括新农人、合作社、农业主管部门和收购/电商企业。", "默认样板为江苏句容场景，数据样板包含台湾农业开放数据与公开农业本体。"]),
                ("核心能力", ["系统不是静态问答页，而是模块化经营辅助平台。每次回答会综合当前页面模块、实时情报、知识库证据和用户输入。"], ["政策解读与申报准备。", "行情与农资查询、价格表循环展示和品类沉淀。", "今天到三周的农业气象风险滑页。", "两地位置、目标市场、路线和沿途散货建议。", "可折叠证据链、知识树和规则沉淀。"]),
                ("交付亮点", ["本次更新加入了无角标头条分类卡、模块按钮深层详情、PPT 新截图和 MP4 演示视频，并打成可信新电脑一键启动包。"], ["解压即用。", "本地知识库和便携运行时随包交付。", "后台模型默认走新农人助手配置。"]),
            ],
        ),
        (
            "03_AgriKB_专利交底书_初稿.docx",
            "03_AgriKB_专利交底书_初稿.md",
            "AgriKB 专利交底书初稿",
            "县域农业多源情报聚合与经营决策辅助方法",
            [
                ("拟保护主题", ["一种面向县域农业经营的多源情报聚合、知识图谱组织和自然语言经营决策辅助方法及系统。"], ["输入包括用户问题、当前页面模块、作物/地块/市场位置、天气、价格、政策、路线和本地知识库。", "输出包括一句话经营建议、政策/市场/气象判断、可执行步骤和可展开证据链。"]),
                ("创新点", ["本系统把可点击模块状态作为问答上下文，避免模型只凭用户一句话生成泛化答案。"], ["模块上下文自动注入检索证据。", "行情查询结果按本体目录沉淀。", "农业气象按今天、明天、三天、十天、十五天、三周滑页化解释。", "路线规划结合市场位置与沿途散货销售建议。"]),
                ("实施方式", ["前端提供经营入口、模块详情、设置、文件管理和联网开关；后端负责检索、实时聚合、模型调用、规则沉淀和本地文件组织。"], ["可部署在普通 Windows 主机、农业局内网电脑、合作社工作站或低功耗边缘终端。", "正式提交前需由专利代理人进一步收敛权利要求。"]),
            ],
        ),
        (
            "04_AgriKB_软件著作权登记材料_初稿.docx",
            "04_AgriKB_软件著作权登记材料_初稿.md",
            "AgriKB 软件著作权登记材料初稿",
            "软件功能、模块和运行环境说明",
            [
                ("软件名称", ["AgriKB 数智新农人经营中枢软件。"], ["建议简称：AgriKB 新农人助手。", "版本号：V1.1。"]),
                ("主要功能模块", ["软件包含智能问答、知识库检索、政策解读、市场行情、农业气象、路线物流、质量溯源、大棚监测、资讯聚合、文件管理、设置管理和便携部署模块。"], ["前端以 HTML/CSS/JavaScript 实现。", "后端以 Python FastAPI 实现。", "本地知识沉淀使用 JSON、JSONL、TTL、CSV 和文档索引组织。"]),
                ("运行环境", ["本交付包包含 Windows 便携 Python 3.11.9、项目虚拟环境包、启动脚本和本地配置文件。"], ["支持 Windows 10/11。", "建议 8GB 内存以上。", "联网功能需要互联网访问；离线仍可使用本地知识库和已有数据。"]),
            ],
        ),
        (
            "05_AgriKB_移动端与硬件部署适配说明.docx",
            "05_AgriKB_移动端与硬件部署适配说明.md",
            "AgriKB 移动端与硬件部署适配说明",
            "可信电脑、手机扫码和农业现场终端部署",
            [
                ("可信电脑一键部署", ["交付包采用自包含目录结构：源码、数据、知识库、便携 Python、依赖包、API 配置、PPT、视频和文档放在同一根目录。"], ["新电脑解压后双击启动。", "一键自检脚本会用独立端口验证 /health 和 UI 静态资源。"]),
                ("手机扫码测试", ["同一局域网内，可使用启动窗口显示的局域网地址或二维码在手机浏览器访问。"], ["若手机无法访问，先确认电脑防火墙和同一 Wi-Fi。", "PWA 页面可添加到手机桌面，适合作为演示和轻量现场应用。"]),
                ("硬件建议", ["农业局或合作社演示可用普通办公电脑；长期部署建议使用小型主机或工控机，外接大屏或触控屏。"], ["可扩展传感器网关。", "可扩展本地地图、电子秤、扫码枪和打印机。"]),
            ],
        ),
    ]
    outputs: list[Path] = []
    for docx_name, md_name, title, subtitle, sections in doc_specs:
        outputs.append(build_docx(docx_name, title, subtitle, sections))
        outputs.append(write_md(md_name, title, sections))
    return outputs


def add_slide_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = ppt_color(PAPER)


def add_textbox(slide, text: str, left, top, width, height, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = ppt_color(color)
    return shape


def add_round_rect(slide, left, top, width, height, fill=WHITE, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ppt_color(fill)
    shape.line.color.rgb = ppt_color(line)
    shape.line.width = PptPt(1)
    return shape


def add_title(slide, kicker: str, title: str, subtitle: str = ""):
    add_textbox(slide, kicker, Inches(0.42), Inches(0.22), Inches(6.8), Inches(0.26), 9.5, GOLD, True)
    add_textbox(slide, title, Inches(0.42), Inches(0.52), Inches(8.8), Inches(0.64), 25, DEEP, True)
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.45), Inches(1.14), Inches(8.8), Inches(0.38), 11.5, MUTED)


def add_image_contain(slide, path: Path, left, top, width, height):
    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(width / iw, height / ih)
    nw, nh = int(iw * scale), int(ih * scale)
    x = left + int((width - nw) / 2)
    y = top + int((height - nh) / 2)
    return slide.shapes.add_picture(str(path), x, y, nw, nh)


def add_screenshot_slide(prs: Presentation, num: int, title_text: str, subtitle: str, image_path: Path, bullets: list[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)
    add_title(slide, f"0{num} / 实测界面", title_text, subtitle)
    add_round_rect(slide, Inches(0.4), Inches(1.55), Inches(6.3), Inches(3.7), WHITE, LINE)
    add_image_contain(slide, image_path, Inches(0.48), Inches(1.65), Inches(6.14), Inches(3.48))
    y = 1.65
    for i, item in enumerate(bullets, 1):
        add_round_rect(slide, Inches(6.95), Inches(y), Inches(2.45), Inches(0.7), "FFFFFF", "D8E4D7")
        add_textbox(slide, f"{i}", Inches(7.08), Inches(y + 0.14), Inches(0.25), Inches(0.25), 14, GREEN, True, PP_ALIGN.CENTER)
        add_textbox(slide, item, Inches(7.38), Inches(y + 0.08), Inches(1.85), Inches(0.46), 10.3, INK, True)
        y += 0.82
    return slide


def build_ppt(video_path: Path, screenshots: list[Path]) -> Path:
    old_decks = list(DELIVERABLES.glob("AgriKB_*路演*.pptx"))
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    shot = {p.name: p for p in screenshots}

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)
    add_round_rect(slide, Inches(0.35), Inches(0.3), Inches(9.3), Inches(5.0), "F9FBF4", "D8E4D7")
    add_textbox(slide, "2026 句容“福地青年英才”创业大赛 · 数智信息 × 新农人", Inches(0.7), Inches(0.75), Inches(8.2), Inches(0.3), 12, GOLD, True)
    add_textbox(slide, "AgriKB\n数智新农人经营中枢", Inches(0.7), Inches(1.12), Inches(5.8), Inches(1.15), 33, DEEP, True)
    add_textbox(slide, "把政策、行情、天气、路线、收购和农业知识变成可执行的经营建议。", Inches(0.72), Inches(2.42), Inches(5.8), Inches(0.42), 15, MUTED)
    for i, (value, label) in enumerate([("一键启动", "可信新电脑解压即用"), ("8 张", "最新实测截图"), ("MP4", "功能讲解视频内嵌"), ("V1.1", "文档与材料同步更新")]):
        x = 0.72 + i * 2.15
        add_round_rect(slide, Inches(x), Inches(3.55), Inches(1.75), Inches(0.74), "EAF4EA", "C8DCC9")
        add_textbox(slide, value, Inches(x + 0.12), Inches(3.67), Inches(1.5), Inches(0.25), 15, GREEN, True, PP_ALIGN.CENTER)
        add_textbox(slide, label, Inches(x + 0.12), Inches(3.95), Inches(1.5), Inches(0.2), 8.5, MUTED, False, PP_ALIGN.CENTER)
    if (shot.get("01_home_clean_v90.png")):
        add_image_contain(slide, shot["01_home_clean_v90.png"], Inches(6.65), Inches(1.1), Inches(2.65), Inches(1.85))
    add_textbox(slide, "路演主张：不是再做一个农业 App，而是把县域农业经营所需的信息、模型和渠道合成一个能落地的助手。", Inches(0.72), Inches(4.62), Inches(8.4), Inches(0.34), 13, INK, True)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)
    add_title(slide, "01 / WHY NOW", "县域农业缺的不是信息，而是可执行判断", "政策、市场、天气、收购、物流分散在不同页面，最终需要回到今天怎么做。")
    for i, (head, body) in enumerate([
        ("新农人", "要知道今天采不采、卖给谁、能不能申报、天气有没有风险。"),
        ("农业局", "要把政策、培训、主体、产业扶持和治理台账连起来。"),
        ("企业", "要看货源、质量、价格、物流和电商渠道是否可靠。"),
    ]):
        add_round_rect(slide, Inches(0.65 + i * 3.05), Inches(1.7), Inches(2.55), Inches(1.35), "FFFFFF", "D8E4D7")
        add_textbox(slide, head, Inches(0.86 + i * 3.05), Inches(1.94), Inches(2.1), Inches(0.28), 17, GREEN if i == 0 else GOLD if i == 1 else RED, True, PP_ALIGN.CENTER)
        add_textbox(slide, body, Inches(0.88 + i * 3.05), Inches(2.35), Inches(2.08), Inches(0.38), 10.6, INK, False, PP_ALIGN.CENTER)
    add_round_rect(slide, Inches(1.0), Inches(3.7), Inches(8.0), Inches(0.7), "EAF4EA", "C8DCC9")
    add_textbox(slide, "AgriKB 的关键价值：把“找信息”变成“下一步怎么干”，把“经验判断”变成“有依据的经营建议”。", Inches(1.18), Inches(3.92), Inches(7.65), Inches(0.24), 14, DEEP, True, PP_ALIGN.CENTER)

    add_screenshot_slide(prs, 2, "产品入口：少操作、强引导", "面向一线用户，入口尽量压缩成几个自然动作。", shot["01_home_clean_v90.png"], ["四个经营入口", "左侧政策/功能栏", "设置与文件入口", "历史对话可清空"])
    add_screenshot_slide(prs, 3, "大棚监测：从指标到动作", "作物本体不直接堆复杂图，而转成因果和调整建议。", shot["02_greenhouse_monitoring_v90.png"], ["实时模拟监测", "达标/偏高一眼可见", "指标可点开", "可继续 AI 分析"])
    add_screenshot_slide(prs, 4, "按钮继续下钻：每个节点都有解释", "用户点到材料、窗口、棚温、市场等节点后，继续看到关联信息。", shot["03_greenhouse_detail_v90.png"], ["节点说明", "所属模块", "下一步动作", "自动聚合上下文"])
    add_screenshot_slide(prs, 5, "资讯聚合：不是新闻列表，而是经营雷达", "政策、天气、行情、收购按农业经营影响重排。", shot["04_information_aggregation_v90.png"], ["聚合公开入口", "头条横向滚动", "可点开来源", "趋势分析"])
    add_screenshot_slide(prs, 6, "大赛主题聚合：六类热点，不展示数量角标", "展开后仍保留每类 10 条实时头条，视觉上更干净。", shot["05_competition_headlines_no_badges_v90.png"], ["农业创业", "科技农业", "青年人才", "政策资金"])
    add_screenshot_slide(prs, 7, "行情与农资：查询后沉淀到本地本体", "价格表动态循环，查询品类会形成可复用本地知识。", shot["06_market_price_loop_v90.png"], ["动态价格表", "手动输入品名", "农机肥料农资", "AI 实时建议"])
    add_screenshot_slide(prs, 8, "天气：从今天到三周的滑页化风险", "默认江苏句容，也支持用户输入城市或地图切换。", shot["07_weather_horizons_v90.png"], ["今天/明天", "三天/十天", "十五天/三周", "生产动作提醒"])
    add_screenshot_slide(prs, 9, "政策雷达：材料、窗口、部门能继续点开", "帮助判断能不能申、先准备什么、找哪个入口。", shot["08_policy_drilldown_v90.png"], ["对象判断", "材料清单", "窗口期", "部门入口"])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)
    add_title(slide, "10 / ARCHITECTURE", "新农人助手会自动聚合当前模块内容", "用户输入不再孤立处理，当前模块、已点开的详情、天气、行情、路线和头条会一起进入后端上下文。")
    layers = [
        ("前端模块", "政策、行情、气象、路线、溯源、大棚监测、资讯聚合"),
        ("上下文聚合", "当前模块 + 点开节点 + 实时面板 + 用户问题"),
        ("检索与模型", "本地知识库、公开数据、联网搜索、新农人助手模型"),
        ("经营回答", "一句话判断、政策机会、市场渠道、气象风险、下一步动作"),
    ]
    for i, (h, b) in enumerate(layers):
        add_round_rect(slide, Inches(0.85 + i * 2.25), Inches(2.0), Inches(1.8), Inches(1.0), "FFFFFF", "D8E4D7")
        add_textbox(slide, h, Inches(0.98 + i * 2.25), Inches(2.2), Inches(1.55), Inches(0.25), 13, GREEN, True, PP_ALIGN.CENTER)
        add_textbox(slide, b, Inches(0.98 + i * 2.25), Inches(2.55), Inches(1.55), Inches(0.24), 8.6, MUTED, False, PP_ALIGN.CENTER)
        if i < 3:
            add_textbox(slide, "→", Inches(2.65 + i * 2.25), Inches(2.32), Inches(0.3), Inches(0.3), 22, GOLD, True, PP_ALIGN.CENTER)
    add_round_rect(slide, Inches(1.0), Inches(3.75), Inches(8.0), Inches(0.72), "EAF4EA", "C8DCC9")
    add_textbox(slide, "结果：回答不再是固定话术，而是围绕当前页面、实时数据和知识库证据组织成清晰的中国式经营建议。", Inches(1.2), Inches(3.98), Inches(7.6), Inches(0.22), 13, DEEP, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)
    add_title(slide, "11 / DELIVERY", "可信新电脑：解压后一键启动", "交付包包含便携 Python、依赖包、知识库、数据、API 配置、PPT、视频和文档。")
    deliver = [
        ("运行包", "runtime\\python311 + .venv\\Lib\\site-packages"),
        ("一键启动", "一键启动AgriKB.bat / Start-TrustedPC.ps1"),
        ("本地配置", ".env 内置新农人助手 API 配置"),
        ("交付材料", "PPT、MP4、说明书、专利交底、软著材料"),
    ]
    for i, (h, b) in enumerate(deliver):
        add_round_rect(slide, Inches(0.75), Inches(1.55 + i * 0.72), Inches(5.55), Inches(0.52), "FFFFFF", "D8E4D7")
        add_textbox(slide, h, Inches(0.95), Inches(1.68 + i * 0.72), Inches(1.2), Inches(0.2), 12, GREEN, True)
        add_textbox(slide, b, Inches(2.0), Inches(1.68 + i * 0.72), Inches(4.0), Inches(0.2), 10.3, INK)
    if QR_PATH.exists():
        add_round_rect(slide, Inches(7.0), Inches(1.6), Inches(1.75), Inches(1.75), "FFFFFF", "D8E4D7")
        add_image_contain(slide, QR_PATH, Inches(7.18), Inches(1.78), Inches(1.38), Inches(1.38))
        add_textbox(slide, "手机扫码测试", Inches(6.85), Inches(3.5), Inches(2.05), Inches(0.25), 12, DEEP, True, PP_ALIGN.CENTER)
    add_textbox(slide, "安全说明：API key 只写入本地 .env，交给可信电脑使用；不要上传到公开 GitHub。", Inches(0.85), Inches(4.75), Inches(8.2), Inches(0.28), 12, RED, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)
    add_title(slide, "12 / BUSINESS", "从比赛项目到县域农业数智服务", "商业路径不是卖单点工具，而是把数据、模型和渠道变成可持续服务。")
    for i, (h, b) in enumerate([
        ("政府端", "政策申报、主体台账、产业监测、助农服务评估。"),
        ("经营端", "合作社/农企订阅行情、收购、溯源和供应链分析。"),
        ("生态端", "电商、物流、农资、培训和新媒体服务入口。"),
    ]):
        add_round_rect(slide, Inches(0.72 + i * 3.05), Inches(1.65), Inches(2.55), Inches(1.5), "FFFFFF", "D8E4D7")
        add_textbox(slide, h, Inches(0.92 + i * 3.05), Inches(1.93), Inches(2.1), Inches(0.25), 16, GREEN if i == 0 else GOLD if i == 1 else RED, True, PP_ALIGN.CENTER)
        add_textbox(slide, b, Inches(0.95 + i * 3.05), Inches(2.37), Inches(2.05), Inches(0.42), 10.2, INK, False, PP_ALIGN.CENTER)
    add_round_rect(slide, Inches(1.0), Inches(3.75), Inches(8.0), Inches(0.8), "EAF4EA", "C8DCC9")
    add_textbox(slide, "路演收束：用江苏句容做县域样板，用台湾农业开放数据做标准化样板，向更多地区复制。", Inches(1.2), Inches(4.02), Inches(7.6), Inches(0.25), 14, DEEP, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide)
    add_title(slide, "13 / DEMO VIDEO", "功能讲解视频已嵌入 PPT", "现场可直接播放，也可在材料包中单独打开 MP4。")
    poster = screenshots[0]
    add_round_rect(slide, Inches(0.8), Inches(1.45), Inches(8.4), Inches(3.65), "FFFFFF", "D8E4D7")
    try:
        slide.shapes.add_movie(str(video_path), Inches(1.0), Inches(1.65), Inches(8.0), Inches(3.25), poster_frame_image=str(poster), mime_type="video/mp4")
    except Exception:
        add_image_contain(slide, poster, Inches(1.0), Inches(1.65), Inches(8.0), Inches(3.25))
        box = add_textbox(slide, f"视频文件：{video_path.name}", Inches(1.3), Inches(4.65), Inches(7.4), Inches(0.22), 12, GREEN, True, PP_ALIGN.CENTER)
        box.click_action.hyperlink.address = str(video_path)

    prs.save(PPTX_PATH)
    for old in old_decks:
        if old.name != PPTX_PATH.name:
            try:
                old.unlink()
            except OSError:
                pass
    return PPTX_PATH


def write_readme(outputs: list[Path], screenshots: list[Path], video_path: Path, pptx_path: Path) -> Path:
    payload = {
        "version": "V1.1",
        "date": TODAY.isoformat(),
        "pptx": str(pptx_path),
        "video": str(video_path),
        "screenshots": [str(p) for p in screenshots],
        "documents": [str(p) for p in outputs],
    }
    (DELIVERABLES / "materials_manifest.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    readme = DELIVERABLES / "README_材料包索引.md"
    lines = [
        "# AgriKB V1.1 交付材料索引",
        "",
        f"- 更新日期：{TODAY.isoformat()}",
        "- 路演 PPT：AgriKB_数智新农人经营中枢_路演终版_20260525.pptx",
        "- 功能讲解视频：AgriKB_function_demo_recording_20260525.mp4",
        "- 最新截图目录：screenshots_20260525/",
        "- 使用说明书、产品说明书、专利交底书、软著材料、移动端与硬件部署说明均已同步更新。",
        "- API key 仅在项目根目录 .env 中作为本地运行配置保存，材料正文不展示明文。",
        "",
        "## 建议演示顺序",
        "",
        "1. 双击一键启动包，打开首页。",
        "2. 演示大棚监测节点下钻。",
        "3. 演示创业大赛专题头条，确认右上角不展示“10条”。",
        "4. 演示行情、气象和政策雷达。",
        "5. 播放 PPT 内嵌功能讲解视频。",
    ]
    readme.write_text("\n".join(lines), encoding="utf-8")
    return readme


def main() -> None:
    ensure_dirs()
    screenshots = copy_latest_screenshots()
    video = make_demo_video(screenshots)
    docs = build_documents()
    pptx = build_ppt(video, screenshots)
    readme = write_readme(docs, screenshots, video, pptx)
    print(json.dumps({
        "pptx": str(pptx),
        "video": str(video),
        "readme": str(readme),
        "documents": [str(p) for p in docs],
        "screenshots": [str(p) for p in screenshots],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
